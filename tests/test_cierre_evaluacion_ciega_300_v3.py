import csv
import hashlib
import json
import re
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data/evaluacion/evaluacion_ciega_300_v3_cerrada_v1"


class TestCierreEvaluacionCiega300V3(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        with (OUT / "referencia_ciega_300_v3.csv").open(encoding="utf-8", newline="") as f: cls.rows = list(csv.DictReader(f))

    def test_cobertura_y_procedencia(self):
        self.assertEqual(len(self.rows), 300)
        self.assertEqual(len({r["intervencion_id"] for r in self.rows}), 300)
        self.assertEqual({r["procedencia"] for r in self.rows}, {"IA asistida; validación humana fila por fila"})
        self.assertEqual(sum(r["incluir_evaluacion"] == "true" for r in self.rows), 299)

    def test_adjudicaciones_humanas(self):
        by_order = {int(r["orden"]): r for r in self.rows}
        self.assertEqual(by_order[1]["etiqueta_v3"], "neutral")
        self.assertEqual(by_order[32]["etiqueta_v3"], "hawkish")
        self.assertEqual(by_order[256]["etiqueta_v3"], "no_puedo_decidir")
        self.assertEqual(by_order[256]["incluir_evaluacion"], "false")

    def test_citas_literales(self):
        norm = lambda x: re.sub(r"\s+", " ", x).strip()
        for r in self.rows:
            if r["es_relevante_v3"] == "1": self.assertIn(norm(r["cita_literal"]), norm(r["texto"]))
            self.assertLessEqual(len(r["cita_literal"]), 300)
        with (OUT / "reparaciones_citas.csv").open(encoding="utf-8", newline="") as f: self.assertEqual(len(list(csv.DictReader(f))), 33)

    def test_workbook_cerrado(self):
        ws = load_workbook(OUT / "evaluacion_ciega_300_v3_cerrada.xlsx", read_only=True, data_only=True)["Anotación"]
        states = [r[11].value for r in ws.iter_rows(min_row=2)]
        self.assertEqual(states.count("cerrado"), 299)
        self.assertEqual(states.count("cerrado_no_decidible"), 1)

    def test_manifest(self):
        manifest = json.loads((OUT / "manifest.json").read_text(encoding="utf-8"))["sha256_salidas"]
        for name, expected in manifest.items(): self.assertEqual(hashlib.sha256((OUT / name).read_bytes()).hexdigest(), expected)


if __name__ == "__main__": unittest.main()
