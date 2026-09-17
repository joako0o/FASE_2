import csv
import hashlib
import json
import unittest
from pathlib import Path

from openpyxl import load_workbook

from scripts.preparar_anotacion_dirigida_300_v3 import OUT, QUOTAS, select, stratum


def read(path):
    with Path(path).open(encoding="utf-8", newline="") as f: return list(csv.DictReader(f))


class TestAnotacionDirigida300V3(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.blind = read(OUT / "anotacion_dirigida_300_v3.csv")
        cls.control = read(OUT / "control_muestreo.csv")
        cls.summary = json.loads((OUT / "resumen.json").read_text(encoding="utf-8"))
        cls.corpus = {r["intervencion_id"]: r for r in read(Path("data/L0/corpus.csv"))}

    def test_cobertura_periodo_y_cuotas(self):
        self.assertEqual(len(self.blind), 300)
        self.assertEqual(len({r["intervencion_id"] for r in self.blind}), 300)
        self.assertEqual(self.summary["cuotas_recuperacion"], QUOTAS)
        counts = {}
        for row in self.control: counts[row["estrato_recuperacion_no_visible"]] = counts.get(row["estrato_recuperacion_no_visible"], 0) + 1
        self.assertEqual(counts, QUOTAS)
        self.assertTrue(all(2005 <= int(self.corpus[r["intervencion_id"]]["anio"]) <= 2015 for r in self.blind))

    def test_instrumento_ciego(self):
        for row in self.blind:
            self.assertEqual((row["etiqueta_v3"], row["es_relevante_v3"], row["confianza"],
                              row["cita_literal"], row["fundamento"], row["estado"]),
                             ("", "", "", "", "", "pendiente"))
        self.assertFalse(self.summary["estrato_visible_en_excel"])
        self.assertFalse(self.summary["etiquetas_visibles"])
        self.assertFalse(self.summary["predicciones_consultadas_o_visibles"])

    def test_excel_no_expone_estrato_y_tiene_validaciones(self):
        wb = load_workbook(OUT / "anotacion_dirigida_300_v3.xlsx", data_only=False)
        self.assertEqual(wb.sheetnames, ["Inicio", "Anotación", "Codebook v3"])
        sheet = wb["Anotación"]
        headers = [sheet.cell(1, col).value for col in range(1, 13)]
        self.assertNotIn("estrato", " ".join(headers).lower())
        self.assertEqual(sheet.max_row, 301)
        self.assertEqual(len(sheet.data_validations.dataValidation), 4)
        self.assertTrue(sheet.protection.sheet)
        self.assertTrue(all(sheet.cell(row, 7).value is None for row in range(2, 302)))

    def test_sin_solapamiento_con_referencias(self):
        selected = {r["intervencion_id"] for r in self.blind}
        for path in ("data/evaluacion/referencia_v3/referencia_v3.csv",
                     "data/evaluacion/anotacion_adicional_300_v3_cerrada_v1/referencia_300_v3.csv"):
            self.assertFalse(selected & {r["intervencion_id"] for r in read(path)})

    def test_texto_fuente_y_estrato_reproducibles(self):
        control = {r["intervencion_id"]: r for r in self.control}
        for row in self.blind:
            source = self.corpus[row["intervencion_id"]]
            self.assertEqual(row["texto"], source["texto"])
            self.assertEqual(hashlib.sha256(row["texto"].encode()).hexdigest(), control[row["intervencion_id"]]["sha256_texto"])
            self.assertEqual(stratum(row["texto"]), control[row["intervencion_id"]]["estrato_recuperacion_no_visible"])
        regenerated, *_ = select()
        self.assertEqual([r["intervencion_id"] for r in regenerated], [r["intervencion_id"] for r in self.blind])

    def test_manifest(self):
        manifest = json.loads((OUT / "manifest.json").read_text(encoding="utf-8"))["sha256_salidas"]
        for name, expected in manifest.items():
            self.assertEqual(hashlib.sha256((OUT / name).read_bytes()).hexdigest(), expected)


if __name__ == "__main__": unittest.main()
