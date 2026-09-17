"""Libro de revisión y descarga estática; sin respuestas IA/humanas privadas."""
from functools import partial
import json
from pathlib import Path
import sys
import tempfile
import threading
import unittest
import urllib.error
import urllib.request

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from utilidades import cargar_script
E=cargar_script('29_preparar_excel_revision.py')


def publico():
    return {'version':'fixture','muestra_sha256':'sha-fixture','guia':'## Guía\nRegla literal.\n',
        'casos':[{'caso_id':f'R{i:02d}','actor':'=1+1' if i==1 else 'Actor de prueba',
                  'cargo':'Consejero','fecha':'2010-01-01',
                  'texto':('=Texto, no fórmula.\n' if i==1 else '')+(f'Intervención inventada {i}. '*140)} for i in range(1,31)]}


def firma(libro):
    return [(h.title,h.freeze_panes,h.protection.sheet,
             [(c.coordinate,c.value,c.data_type,c.protection.locked,c.hyperlink.location if c.hyperlink else None)
              for fila in h for c in fila],
             [(str(d.sqref),d.type,d.formula1,d.showErrorMessage) for d in h.data_validations.dataValidation]) for h in libro]


class ExcelRevision(unittest.TestCase):
    def test_bloques_sin_perdidas_incluso_saltos_y_palabras_largas(self):
        for s in ['texto corto',' '*1600,'a'*2200,'\n'*600,'abc\n'*200,'Última frase con inflación. '*100]:
            partes=E.bloques(s)
            self.assertEqual(''.join(partes),s)
            self.assertTrue(all(len(p)<=650 for p in partes))
            self.assertTrue(all(p.count('\n')<=8 for p in partes))

    def test_libro_completo_celdas_vacias_y_reproducible(self):
        with tempfile.TemporaryDirectory() as t:
            p=Path(t)/'primero.xlsx';q=Path(t)/'segundo.xlsx'
            datos=publico()
            ubicaciones=E.crear_libro(datos,p)
            E.crear_libro(datos,q)
            a,b=load_workbook(p),load_workbook(q)
            self.assertEqual(firma(a),firma(b))
            self.assertEqual(a.sheetnames,['Inicio','Respuestas','Textos','Guía'])
            self.assertTrue(all(h.sheet_state=='visible' for h in a))
            self.assertEqual(a['Inicio']['B22'].value,datos['muestra_sha256'])
            self.assertEqual(a['Inicio']['B21'].value,datos['version'])
            self.assertEqual(a['Respuestas'].max_row,36)
            for caso in datos['casos']:
                u=ubicaciones[caso['caso_id']];fila=u['fila_respuesta']
                self.assertEqual(a['Respuestas'].cell(fila,1).value,caso['caso_id'])
                self.assertEqual(''.join(a['Textos'].cell(i,3).value for i in u['filas_texto']),caso['texto'])
                self.assertTrue(all(a['Textos'].row_dimensions[i].height<=409 for i in u['filas_texto']))
                for col in range(3,7):
                    c=a['Respuestas'].cell(fila,col)
                    self.assertIsNone(c.value)
                    self.assertFalse(c.protection.locked)
                self.assertEqual(a['Respuestas'].cell(fila,7).value,E.formula_estado(fila))
                self.assertTrue(a['Respuestas'].cell(fila,7).protection.locked)
                self.assertTrue(a['Respuestas'].cell(fila,8).hyperlink.location.startswith("'Textos'!"))
            self.assertEqual(a['Respuestas']['B7'].value,'=1+1')
            self.assertEqual(a['Respuestas']['B7'].data_type,'s')
            formulas=[(h.title,c.coordinate) for h in a for fila in h for c in fila if c.data_type=='f']
            self.assertEqual(len(formulas),31)
            self.assertFalse(a._external_links)
            self.assertEqual(len(a['Respuestas'].data_validations.dataValidation),3)
            a.close();b.close()
            with self.assertRaises(FileExistsError): E.crear_libro(datos,p)

    def test_formula_campos_no_es_puntuacion_ia(self):
        f=E.formula_estado(7)
        self.assertEqual(f.count('('),f.count(')'))
        self.assertTrue(Tokenizer(f).items)
        self.assertIn('C7="No puedo decidir"',f)
        self.assertIn('LEN(E7)>300',f)
        self.assertNotIn('Textos!',f)  # No simular validación literal que no se realiza.
        self.assertNotIn('clave',f.lower())

    def test_enlace_http_no_modifica_payload_o_javascript(self):
        plantilla=(E.ANTERIOR.RAIZ/'scripts/plantillas/revision_entrenamiento.html').read_text()
        original=plantilla.replace('__DATOS_REVISION__',json.dumps(publico()))
        nuevo=E.actualizar_html(original)
        self.assertEqual(E.ANTERIOR.payload(original),E.ANTERIOR.payload(nuevo))
        self.assertEqual(original.split('<script>')[1],nuevo.split('<script>')[1])
        self.assertIn(f'href="{E.NOMBRE_EXCEL}"',nuevo)
        with self.assertRaises(AssertionError): E.actualizar_html(nuevo)

    def test_descarga_es_xlsx_real_y_rutas_privadas_no_se_exponen(self):
        with tempfile.TemporaryDirectory() as t:
            base=Path(t);(base/'formulario').mkdir()
            (base/'formulario/index.html').write_text('<html>Revisión</html>')
            E.crear_libro(publico(),base/E.NOMBRE_EXCEL)
            (base/'NO_ABRIR_hasta_finalizar_clave.csv').write_text('no público')
            servidor=E.ThreadingHTTPServer(('0.0.0.0',0),partial(E.DescargaRevision,directorio=base))
            hilo=threading.Thread(target=servidor.serve_forever,daemon=True);hilo.start()
            raiz=f'http://127.0.0.1:{servidor.server_port}'
            try:
                with urllib.request.urlopen(raiz+'/'+E.NOMBRE_EXCEL) as r:
                    self.assertEqual(r.status,200)
                    self.assertIn('attachment',r.headers['Content-Disposition'])
                    self.assertIn('spreadsheetml.sheet',r.headers['Content-Type'])
                    self.assertEqual(r.read(),(base/E.NOMBRE_EXCEL).read_bytes())
                with urllib.request.urlopen(urllib.request.Request(raiz+'/'+E.NOMBRE_EXCEL,method='HEAD')) as r:
                    self.assertEqual(r.status,200);self.assertEqual(r.read(),b'')
                for ruta in ['/NO_ABRIR_hasta_finalizar_clave.csv','/../protocolo.json','/%2e%2e/manifest.json','/formulario/','/scripts/29_preparar_excel_revision.py']:
                    with self.assertRaises(urllib.error.HTTPError) as e: urllib.request.urlopen(raiz+ruta)
                    self.assertEqual(e.exception.code,404)
            finally:
                servidor.shutdown();servidor.server_close();hilo.join()


if __name__=='__main__': unittest.main()
