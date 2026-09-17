"""Fixtures sintéticos de software, nunca ejemplos para entrenar ni respuestas del gold."""
import copy
import hashlib
import json
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
import muestreo_revision as M


class MuestreoRevision(unittest.TestCase):
    def corpus(self):
        filas=[]
        for canal in ['H','D']:
            for tipo in ['decision','trayectoria','contraste']:
                for variante in range(3):
                    k=f'{canal}_{tipo}_{variante}'
                    verbo='subir' if canal=='H' else 'bajar'
                    prefijo={'decision':'Mi voto es por ','trayectoria':'Se prevé ','contraste':'Sin embargo se prevé '}[tipo]
                    texto=prefijo+verbo+' la TPM. '+ ' '.join(f'palabra{k}{i}' for i in range(35))
                    filas.append({'intervencion_id':k,'meeting_id':k,'actor':k,'anio':str(2005+variante),
                        'texto':texto,'fecha':'2005-01-01','cargo':'fixture','flag_texto_danado':'False'})
        return filas

    def test_01_pistas_no_son_etiquetas(self):
        self.assertIsNone(M.pistas('La inflación sube y el desempleo baja.'))
        self.assertEqual(M.pistas('Se propone reducir el estímulo monetario.')['canal_busqueda'],'H')
        self.assertEqual(M.pistas('Se considera mayor estímulo monetario.')['canal_busqueda'],'D')
        self.assertIsNotNone(M.pistas('No es conveniente subir la TPM.'))  # Detecta mención, no adhesión.

    def test_02_seleccion_reproducible_y_cuotas(self):
        c=self.corpus();a,r=M.seleccionar(c,set(),set(),por_celda=1)
        b,_=M.seleccionar(list(reversed(c)),set(),set(),por_celda=1)
        self.assertEqual([x['intervencion_id'] for x in a],[x['intervencion_id'] for x in b])
        self.assertEqual(len(a),6);self.assertEqual(sum(x['canal_busqueda']=='H' for x in a),3)
        self.assertEqual(r['solapes_ids'],0)

    def test_03_excluye_ids_y_copias_normalizadas(self):
        c=self.corpus();k=c[0]['intervencion_id'];copia={**c[0],'intervencion_id':'copia','texto':'  '+c[0]['texto'].upper()+'\n'}
        a,r=M.seleccionar(c+[copia],{k},{c[-1]['intervencion_id']},por_celda=1)
        self.assertFalse({x['intervencion_id'] for x in a}&{k,'copia',c[-1]['intervencion_id']})
        self.assertEqual(r['solapes_textos_normalizados'],0)

    def test_04_danados_y_casi_copias(self):
        c=self.corpus();c[0]['flag_texto_danado']='True'
        a,_=M.seleccionar(c,set(),set(),por_celda=1)
        self.assertNotIn(c[0]['intervencion_id'],{x['intervencion_id'] for x in a})
        f=M.trigramas(' '.join('palabra'+str(i) for i in range(100)))
        self.assertTrue(M.casi_copia(f,f));self.assertFalse(M.casi_copia(f,M.trigramas('un ejemplo totalmente ajeno')))

    def test_05_fallo_ante_pool_insuficiente(self):
        with self.assertRaises(ValueError):M.seleccionar(self.corpus()[:1],set(),set())
        c=self.corpus();c.append(c[0])
        with self.assertRaises(ValueError):M.seleccionar(c,set(),set())

    def test_06_purga_por_reunion_no_cambia_folds(self):
        c=self.corpus();c[1]['meeting_id']=c[0]['meeting_id']
        folds=[{'fold':1,'train':[c[2]['intervencion_id']],'validacion':[c[0]['intervencion_id']]},
               {'fold':2,'train':[],'validacion':[c[-1]['intervencion_id']]}]
        anterior=copy.deepcopy(folds);r=M.plan_folds([c[1]],c,folds)[c[1]['intervencion_id']]
        self.assertEqual(r['folds_prohibidos'],[1]);self.assertEqual(r['folds_potencialmente_permitidos'],[2])
        self.assertFalse(r['incorporado_a_train']);self.assertEqual(folds,anterior)

    def test_07_excel_sin_claves_textos_completos_y_respuestas_vacias(self):
        from openpyxl import load_workbook
        casos=[]
        for i in range(1,61):
            casos.append({'caso_id':f'C{i:02d}','bloque':(i-1)//10+1,'actor':'=NO_EJECUTAR()',
                'cargo':'fixture','fecha':'2005-01-01','texto':'=NO_EJECUTAR()\n'+('texto\n'*100 if i==1 else 'palabras '*130),
                'canal_busqueda':'SECRETO_DE_SELECCION'})
        with tempfile.TemporaryDirectory() as t:
            p=Path(t)/'revision.xlsx';u=M.crear_excel(casos,'Guía sin etiquetas por caso',p)
            wb=load_workbook(p)
            self.assertEqual(len(wb.sheetnames),14)
            for r in casos:
                loc=u[r['caso_id']];h=wb[loc['hoja_respuesta']];f=loc['fila_respuesta']
                for col in range(3,7):self.assertIsNone(h.cell(f,col).value);self.assertFalse(h.cell(f,col).protection.locked)
                self.assertEqual(h.cell(f,2).data_type,'s')
                reconstruido=''.join(wb[loc['hoja_texto']].cell(n,3).value for n in loc['filas_texto'])
                self.assertEqual(reconstruido,r['texto'])
                self.assertTrue(h.cell(f,8).hyperlink.location)
            self.assertTrue(all(h.sheet_state=='visible' for h in wb))
            self.assertFalse(wb._external_links)
            with zipfile.ZipFile(p) as z:
                self.assertFalse(any(b'SECRETO_DE_SELECCION' in z.read(n) for n in z.namelist()))
                self.assertFalse(any('vbaProject' in n for n in z.namelist()))
            with self.assertRaises(FileExistsError):M.crear_excel(casos,'Guía',p)

    def test_08_salida_existente_no_se_sobrescribe(self):
        with tempfile.TemporaryDirectory() as t:
            with self.assertRaises(FileExistsError):M.ejecutar(Path(t))


if __name__=='__main__':unittest.main()
