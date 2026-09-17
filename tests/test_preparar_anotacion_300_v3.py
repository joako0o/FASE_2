import csv
import hashlib
import json
import subprocess
import tempfile
import unittest
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data/muestras/anotacion_adicional_300_v3"
CSV_PATH = OUT / "anotacion_adicional_300_v3.csv"
XLSX_PATH = OUT / "anotacion_adicional_300_v3.xlsx"
SCRIPT = ROOT / "scripts/preparar_anotacion_300_v3.py"
PREVIAS = [
    "estrato_enriquecido.csv", "estrato_fases.csv", "estrato_tanda9.csv",
    "gold_ciego_300.csv", "piloto_300.csv", "test_retest_30.csv",
]
RESPUESTAS = ["etiqueta_v3", "es_relevante_v3", "confianza", "cita_literal", "fundamento"]


def leer(path):
    with path.open(encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


class TestPrepararAnotacion300V3(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.csv_antes = CSV_PATH.read_bytes()
        cls.tmp = tempfile.TemporaryDirectory()
        salida = Path(cls.tmp.name) / "repeticion"
        subprocess.run(
            [str(ROOT / ".venv/bin/python"), str(SCRIPT), "--salida", str(salida)],
            cwd=ROOT, check=True, capture_output=True, text=True,
        )
        cls.csv_repetido = (salida / CSV_PATH.name).read_bytes()
        cls.rows = leer(CSV_PATH)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_reproducible(self):
        self.assertEqual(self.csv_antes, self.csv_repetido)

    def test_unicidad_topes_y_campos_ciegos(self):
        self.assertEqual(300, len(self.rows))
        ids = [x["intervencion_id"] for x in self.rows]
        self.assertEqual(300, len(set(ids)))
        corpus = {x["intervencion_id"]: x for x in leer(ROOT / "data/L0/corpus.csv")}
        self.assertLessEqual(max(Counter(corpus[x]["meeting_id"] for x in ids).values()), 4)
        self.assertLessEqual(max(Counter(corpus[x]["anio"] for x in ids).values()), 32)
        for row in self.rows:
            self.assertTrue(row["texto"].strip())
            self.assertTrue(row["actor"].strip())
            self.assertTrue(row["cargo"].strip())
            self.assertTrue(row["fecha_reunion"].strip())
            self.assertTrue(all(row[c] == "" for c in RESPUESTAS))
        encabezados = {k.lower() for k in self.rows[0]}
        self.assertFalse(any("pred" in k or "etiqueta_previa" in k for k in encabezados))

    def test_texto_fiel_y_exclusiones(self):
        corpus = {x["intervencion_id"]: x for x in leer(ROOT / "data/L0/corpus.csv")}
        sample_ids = {x["intervencion_id"] for x in self.rows}
        for row in self.rows:
            self.assertEqual(corpus[row["intervencion_id"]]["texto"], row["texto"])
        ref_ids = {x["intervencion_id"] for x in leer(ROOT / "data/evaluacion/referencia_v3/referencia_v3.csv")}
        self.assertFalse(sample_ids & ref_ids)
        for nombre in PREVIAS:
            prev_ids = {x["intervencion_id"] for x in leer(ROOT / "data/muestras" / nombre)}
            self.assertFalse(sample_ids & prev_ids, nombre)

    def test_manifest_integridad(self):
        manifest = json.loads((OUT / "manifest.json").read_text(encoding="utf-8"))
        hashes = manifest["sha256_salidas"]
        self.assertEqual(sha(CSV_PATH), hashes[CSV_PATH.name])
        self.assertEqual(sha(XLSX_PATH), hashes[XLSX_PATH.name])

    def test_workbook_schema_validaciones_y_formulas(self):
        wb = load_workbook(XLSX_PATH, data_only=False)
        self.assertEqual(["Inicio", "Anotación", "Codebook v3"], wb.sheetnames)
        ws = wb["Anotación"]
        headers = [c.value for c in ws[1]]
        self.assertEqual(list(self.rows[0]), headers)
        self.assertEqual(301, ws.max_row)
        self.assertTrue(ws.protection.sheet)
        self.assertGreaterEqual(len(ws.data_validations.dataValidation), 4)
        cols = {name: i + 1 for i, name in enumerate(headers)}
        for r in range(2, 302):
            self.assertIsNone(ws.cell(r, cols["etiqueta_v3"]).value)
            self.assertIsNone(ws.cell(r, cols["es_relevante_v3"]).value)
            self.assertIsNone(ws.cell(r, cols["confianza"]).value)
            self.assertIsNone(ws.cell(r, cols["cita_literal"]).value)
            self.assertIsNone(ws.cell(r, cols["fundamento"]).value)
            self.assertTrue(str(ws.cell(r, cols["estado"]).value).startswith("=IF("))
            self.assertTrue(ws.cell(r, cols["texto"]).protection.locked)
            self.assertFalse(ws.cell(r, cols["etiqueta_v3"]).protection.locked)


if __name__ == "__main__":
    unittest.main()
