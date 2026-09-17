import csv
import hashlib
import json
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from scripts.preparar_evaluacion_ciega_300_v3 import INSTRUMENT_PATHS, OUT, REFERENCE_PATHS, select


def read(path):
    with Path(path).open(encoding="utf-8", newline="") as f: return list(csv.DictReader(f))


class TestEvaluacionCiega300V3(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.rows = read(OUT / "evaluacion_ciega_300_v3.csv"); cls.meetings = pd.read_csv(OUT / "reuniones_cuarentena.csv", dtype=str)
        cls.summary = json.loads((OUT / "resumen.json").read_text(encoding="utf-8"))

    def test_cobertura_y_grupos(self):
        self.assertEqual(len(self.rows), 300); self.assertEqual(len({r["intervencion_id"] for r in self.rows}), 300)
        self.assertEqual(self.meetings.meeting_id.nunique(), 33); self.assertEqual(self.meetings.groupby("anio").meeting_id.nunique().to_dict(), {str(y): 3 for y in range(2005, 2016)})
        quarantine = set(self.meetings.meeting_id)
        self.assertTrue(all(r["intervencion_id"].rsplit(":", 2)[0] in quarantine for r in self.rows))

    def test_ciega_y_sin_predicciones(self):
        for row in self.rows:
            self.assertEqual((row["etiqueta_v3"], row["es_relevante_v3"], row["confianza"], row["cita_literal"], row["fundamento"], row["estado"]), ("", "", "", "", "", "pendiente"))
        self.assertFalse(self.summary["etiquetas_visibles"]); self.assertFalse(self.summary["predicciones_visibles_o_consultadas"])

    def test_sin_solapamiento_previo(self):
        selected = {r["intervencion_id"] for r in self.rows}
        for path in REFERENCE_PATHS + INSTRUMENT_PATHS: self.assertFalse(selected & set(pd.read_csv(path, dtype=str, keep_default_na=False).intervencion_id))

    def test_estratos_ocultos_y_reproducibilidad(self):
        self.assertEqual(self.summary["estratos_recuperacion_ocultos"], {"aleatorio_contextual": 109, "candidato_dovish": 41, "candidato_hawkish": 71, "neutral_dificil": 79})
        selected, meetings, _ = select()
        self.assertEqual(selected.intervencion_id.tolist(), [r["intervencion_id"] for r in self.rows])
        self.assertEqual(set(meetings.meeting_id), set(self.meetings.meeting_id))

    def test_excel_protegido_no_expone_estrato(self):
        wb = load_workbook(OUT / "evaluacion_ciega_300_v3.xlsx", data_only=False)
        self.assertEqual(wb.sheetnames, ["Inicio", "Anotación", "Codebook v3"]); ws = wb["Anotación"]
        self.assertEqual(ws.max_row, 301); self.assertTrue(ws.protection.sheet); self.assertEqual(len(ws.data_validations.dataValidation), 4)
        self.assertNotIn("estrato", " ".join(str(ws.cell(1, col).value) for col in range(1, 13)).lower())

    def test_manifest(self):
        manifest = json.loads((OUT / "manifest.json").read_text(encoding="utf-8"))["sha256_salidas"]
        for name, expected in manifest.items(): self.assertEqual(hashlib.sha256((OUT / name).read_bytes()).hexdigest(), expected)


if __name__ == "__main__": unittest.main()
