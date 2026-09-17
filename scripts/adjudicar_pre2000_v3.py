"""Consolida la adjudicación 257/257 del set pre-2000 bajo codebook v3."""
import csv
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Set_Entrenamiento_Pre_2000.xlsx"
OUT = ROOT / "data/evaluacion/pre2000_adjudicacion_v3_v1"

# Correcciones semánticas detectadas en lectura completa. Las demás filas se ratifican.
CHANGES = {
    21889: ("dovish", "1", "media", "El pronunciamiento hace propia y defiende una trayectoria de disminuciones paulatinas de tasas como política prudente; por tanto respalda dirección expansiva, aunque rechace acelerar los recortes."),
    21910: ("neutral", "1", "alta", "Relata rebajas y sus efectos pasados sin adoptar, recomendar ni respaldar una nueva dirección en el pronunciamiento actual."),
    27869: ("", "1", "media", "La unidad fundamenta «esta decisión», pero la dirección de esa decisión quedó fuera del fragmento. No se hereda del fragmento vecino; queda pendiente."),
    29596: ("neutral", "1", "alta", "Fija niveles de facilidades y mezcla normas cambiarias, pero no expresa comparación previa ni orientación monetaria respaldada."),
    31369: ("neutral", "1", "alta", "Presenta mantener o relajar y dos magnitudes de baja sin resolver entre mantener y bajar; un menú sin preferencia es N."),
    33261: ("neutral", "0", "alta", "Adopta una expectativa sobre la tasa del FED, no una orientación monetaria doméstica para Chile."),
    33267: ("neutral", "1", "alta", "Formula preguntas y describe expectativas de mercado; no adopta una recomendación propia sobre la tasa chilena."),
    33546: ("neutral", "1", "alta", "Diagnostica mayores expectativas de inflación y recomienda disciplina fiscal, sin conectar ese diagnóstico con una dirección monetaria propia."),
    21646: ("neutral", "1", "alta", "Fija niveles de facilidades sin indicar nivel anterior, cambio ni sesgo; el nivel absoluto no determina H/D."),
    21715: ("neutral", "1", "alta", "Fija el Depósito de Liquidez en 5% sin indicar cambio o dirección dentro de la unidad; no se hereda el nivel previo de otra intervención."),
    22418: ("neutral", "1", "alta", "Reduce un porcentaje reglamentario de acceso a una línea, pero la unidad no documenta su función como orientación monetaria agregada; el verbo reducir no basta."),
    30714: ("neutral", "1", "alta", "Fija niveles de la Línea de Crédito de Liquidez sin comparación previa ni dirección respaldada."),
    31398: ("neutral", "1", "alta", "Pide que el Banco actúe ante menor crecimiento, pero no identifica qué acción o dirección monetaria respalda; el diagnóstico no asigna D."),
    31427: ("neutral", "1", "alta", "Menciona una reacción coherente ante riesgo contractivo sin especificar una acción o dirección monetaria; no se infiere una baja."),
    32872: ("neutral", "0", "alta", "Es una sanción operativa dirigida a instituciones con déficit, no una orientación monetaria agregada documentada."),
    32888: ("neutral", "0", "alta", "Reduce un encaje a influjos de capital como regulación cambiaria/financiera; no documenta una dirección de política monetaria doméstica."),
    33274: ("neutral", "1", "alta", "Descarta una acción monetaria macroeconómica no especificada; no puede invertirse ese rechazo ni suponerse que la acción descartada era contractiva."),
    33524: ("neutral", "1", "alta", "Relata una decisión previa sobre la meta de inflación y no adopta una dirección monetaria actual; evitar afectar la evolución no equivale por sí solo a D."),
}
PENDING_ORIGINAL = {31972, 22077, 32181, 33268, 30995, 31991}


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def norm(text):
    return " ".join(str(text or "").split())


def literal_quote(quote, text):
    q, t = norm(quote), norm(text)
    if len(q) <= 300 and q in t:
        return q
    # Las ocho incidencias originales terminan en elipsis añadida a un prefijo literal.
    q = q.removesuffix("...").rstrip()
    if len(q) > 300:
        q = q[:300].rstrip()
    while q and q not in t:
        q = q[:-1].rstrip()
    if not q:
        raise ValueError("No fue posible recuperar cita literal")
    return q


def run(out=OUT):
    out = Path(out)
    if out.exists():
        raise FileExistsError("No sobrescribir adjudicación")
    ws = load_workbook(SOURCE, read_only=True, data_only=True).active
    headers = [c.value for c in ws[1]]
    rows = [dict(zip(headers, values)) for values in ws.iter_rows(min_row=2, values_only=True)]
    if len(rows) != 257 or len({int(r["frag_id"]) for r in rows}) != 257:
        raise ValueError("Fuente pre-2000 inesperada")

    output = []
    for row in rows:
        fid = int(row["frag_id"])
        old = str(row["etiqueta"])
        if fid in CHANGES:
            label, relevance, confidence, rationale = CHANGES[fid]
            state = "pendiente" if not label else "cerrado"
            result = "corregido_v3"
        elif fid in PENDING_ORIGINAL:
            label, relevance, confidence = "", str(row["es_relevante"]), str(row["confianza"])
            rationale = str(row["nota"])
            state, result = "pendiente", "pendiente_preservado"
        else:
            label, relevance, confidence = old, str(row["es_relevante"]), str(row["confianza"])
            rationale = str(row["nota"])
            state, result = "cerrado", "ratificado_v3"
        quote = literal_quote(row["frase_justificante"], row["texto_original"])
        output.append({
            "frag_id": fid, "fecha": row["fecha"], "anio": row["anio"],
            "etiqueta_origen": old, "es_relevante_origen": row["es_relevante"],
            "etiqueta_v3": label, "es_relevante_v3": relevance,
            "confianza_revision_v3": confidence, "cita_literal": quote,
            "fundamento_revision_v3": rationale, "estado_revision_v3": state,
            "resultado_revision": result, "tipo_accion": row["tipo_accion"],
            "instrumento": row["instrumento"], "texto_original": row["texto_original"],
            "sha256_texto": hashlib.sha256(norm(row["texto_original"]).encode()).hexdigest(),
        })

    out.mkdir(parents=True)
    csv_path = out / "adjudicacion_pre2000_v3.csv"
    with csv_path.open("x", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=output[0], lineterminator="\n")
        writer.writeheader(); writer.writerows(output)
    counts = lambda key: {v: sum(r[key] == v for r in output) for v in sorted({r[key] for r in output})}
    summary = {
        "filas": 257, "ids_unicos": 257, "ratificadas": counts("resultado_revision").get("ratificado_v3", 0),
        "corregidas": counts("resultado_revision").get("corregido_v3", 0),
        "pendientes_preservadas": counts("resultado_revision").get("pendiente_preservado", 0),
        "etiquetas_v3": counts("etiqueta_v3"), "relevancia_v3": counts("es_relevante_v3"),
        "estados": counts("estado_revision_v3"), "incorporacion_entrenamiento": "solo estado_revision_v3=cerrado",
    }
    summary_path = out / "resumen.json"
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    protocol = {
        "codebook": "docs/codebook_v3.md", "source": str(SOURCE.relative_to(ROOT)),
        "source_sha256": sha(SOURCE), "scope": "adjudicacion semantica 257/257",
        "rules_checked": ["direccion propia respaldada", "no contexto invisible", "no signo por verbo o nivel", "objeto domestico", "cita literal continua <=300"],
        "pending_excluded_from_training": True,
    }
    protocol_path = out / "protocolo.json"
    protocol_path.write_text(json.dumps(protocol, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    manifest = {"sha256_salidas": {p.name: sha(p) for p in [csv_path, summary_path, protocol_path]}}
    (out / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return summary


if __name__ == "__main__":
    print(json.dumps(run(), ensure_ascii=False, indent=2))
