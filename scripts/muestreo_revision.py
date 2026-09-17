"""Selección reproducible de candidatos reales; no etiqueta, sintetiza ni entrena.

Entrada para el usuario: 40_gestionar_proyecto.py preparar-muestra-hd.
Reutiliza utilidades Excel existentes sin modificar etapas congeladas.
"""
# ---- 1. Parámetros y búsqueda lexical: canales, no clases verdaderas ----
import argparse
from collections import Counter
from datetime import datetime
import hashlib
import json
import math
from pathlib import Path
import re
import unicodedata

import pandas as pd
from utilidades import cargar_script, exigir_salidas_nuevas, sha256

RAIZ = Path(__file__).resolve().parents[1]
VERSION = 'ampliacion_hd_60_v1'
SEMILLA = 20260916
SALIDA = RAIZ/'data/auditoria'/VERSION
NOMBRE = 'revision_60_candidatos.xlsx'
PATRONES = {
    'H': [r'(?:subir|elevar|aumentar|incrementar|alza de|aumento de).{0,65}(?:tpm|tasa de (?:politica|interes))',
          r'(?:retiro|retirar|reduccion|reducir).{0,30}(?:estimulo|impulso) monetario',
          r'sesgo al alza'],
    'D': [r'(?:bajar|rebajar|reducir|recortar|baja de|recorte de|reduccion de).{0,65}(?:tpm|tasa de (?:politica|interes))',
          r'(?:mayor|aumentar|mantener|mas|incrementar).{0,30}(?:estimulo|impulso) monetario',
          r'sesgo (?:expansivo|a la baja)'],
}
DECISION = r'(?:voto|vota|acordo|decidio|mi opcion|su opcion|su preferencia).{0,100}(?:subir|bajar|elevar|reducir|recortar|aumentar|rebajar)'
CONTRASTE = r'\b(?:sin embargo|no obstante|pausa|prematuro|aunque|no descarta)\b'


def clave(texto):
    return ''.join(c for c in unicodedata.normalize('NFKD', ' '.join(texto.lower().split()))
                   if not unicodedata.combining(c))


def rango(identificador, etapa):
    return hashlib.sha256(f'{SEMILLA}|{etapa}|{identificador}'.encode()).hexdigest()


def pistas(texto):
    t = clave(texto)
    # No cruza oraciones con punto; buscar no equivale a determinar adhesión/país.
    oraciones = re.split(r'[.!?;]', t)
    scores = {c:sum(bool(re.search(p, o)) for p in patrones for o in oraciones)
              for c, patrones in PATRONES.items()}
    if not any(scores.values()):
        return None
    canal = max(scores, key=lambda c:(scores[c], rango(t[:80], c)))
    tipo = ('contraste' if re.search(CONTRASTE, t) or all(scores.values()) else
            'decision' if re.search(DECISION, t) else 'trayectoria')
    return {'canal_busqueda':canal, 'tipo_busqueda':tipo, 'coincidencias':scores}


# ---- 2. Exclusiones por ID/texto y diversidad; nunca leer respuestas del gold ----
def cargar(raiz=RAIZ):
    raiz = Path(raiz)
    fuentes = [raiz/'data/L0/corpus.csv', raiz/'data/muestras/gold_ciego_300.csv',
               raiz/'data/checkpoints/beto_v1/entrada/documentos.json',
               raiz/'data/checkpoints/beto_v1/entrada/folds.json']
    corpus = pd.read_csv(fuentes[0], dtype=str, keep_default_na=False).to_dict('records')
    humanos = set(pd.read_csv(fuentes[1], usecols=['intervencion_id'], dtype=str).intervencion_id)
    anotados = {d['intervencion_id'] for d in json.loads(fuentes[2].read_text(encoding='utf-8'))}
    for p in sorted((raiz/'data/etiquetas').glob('*.csv')):
        fuentes.append(p)
        anotados.update(pd.read_csv(p, usecols=['intervencion_id'], dtype=str).intervencion_id)
    if len(humanos) != 306 or len(anotados) != 1352:
        raise ValueError('Cambió el marco de exclusiones; revisar antes de generar otra muestra')
    folds = json.loads(fuentes[3].read_text(encoding='utf-8'))
    return corpus, anotados, humanos, folds, {str(p.relative_to(raiz)):sha256(p) for p in fuentes}


def trigramas(texto):
    palabras=re.findall(r"\w+", clave(texto))
    return set(zip(palabras,palabras[1:],palabras[2:]))


def casi_copia(a,b,umbral=.85):
    if not a or not b or min(len(a),len(b))/max(len(a),len(b)) < umbral: return False
    inter=len(a & b)
    return inter/(len(a)+len(b)-inter) >= umbral


def seleccionar(corpus, anotados, humanos, por_celda=10):
    ids = [r['intervencion_id'] for r in corpus]
    if len(ids) != len(set(ids)) or not (anotados|humanos) <= set(ids):
        raise ValueError('IDs ausentes o duplicados')
    vetados = {clave(r['texto']) for r in corpus if r['intervencion_id'] in anotados|humanos}
    pool = []; descartes = Counter(); vistos = set(vetados)
    for r in sorted(corpus, key=lambda r:rango(r['intervencion_id'], 'deduplicar')):
        k = clave(r['texto'])
        if r['intervencion_id'] in anotados|humanos or k in vistos:
            descartes['id_o_texto_ya_visto'] += 1; continue
        if r.get('flag_texto_danado', '').lower() == 'true' or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', r['texto']):
            descartes['danado_o_xml_invalido'] += 1; continue
        if len(r['texto']) < 200:
            descartes['menos_de_200_caracteres'] += 1; continue
        p = pistas(r['texto'])
        if p is None:
            descartes['sin_pistas'] += 1; continue
        vistos.add(k)
        pool.append({**r, **p})
    elegidos = []; reuniones = Counter(); actores = Counter(); anios = Counter(); longitudes = Counter()
    def longitud(r):
        return 'corta' if len(r['texto']) <= 1800 else 'media' if len(r['texto']) <= 4000 else 'larga'
    celdas = [(c, t) for c in ['H','D'] for t in ['decision','trayectoria','contraste']]
    conteos = Counter((r['canal_busqueda'], r['tipo_busqueda']) for r in pool)
    # Se preserva la cuota de canal. Si una familia es escasa se registra la reasignación.
    cuotas = {celda:min(por_celda, conteos[celda]) for celda in celdas}
    for canal in ['H','D']:
        faltan = 3*por_celda-sum(n for (c,t),n in cuotas.items() if c == canal)
        for tipo in ['contraste','decision','trayectoria']:
            celda=(canal,tipo);extra=min(faltan, conteos[celda]-cuotas[celda]);cuotas[celda]+=extra;faltan-=extra
        if faltan: raise ValueError('No hay suficientes candidatos para el canal '+canal)
    usados = set(); rechazados_copia=set()
    firmas={r['intervencion_id']:trigramas(r['texto']) for r in pool}
    firmas_vetadas=[trigramas(r['texto']) for r in corpus if r['intervencion_id'] in anotados|humanos]
    while len(elegidos) < 6*por_celda:
        progreso = False
        for canal,tipo in celdas:
            if sum(r['canal_busqueda']==canal and r['tipo_busqueda']==tipo for r in elegidos) >= cuotas[(canal,tipo)]: continue
            candidatos = [r for r in pool if r['canal_busqueda']==canal and r['tipo_busqueda']==tipo and
                r['intervencion_id'] not in usados|rechazados_copia and reuniones[r['meeting_id']] < 2 and actores[r['actor']] < 12]
            if not candidatos: raise ValueError('No se cumplen cuotas/diversidad; no relajar silenciosamente')
            r = min(candidatos, key=lambda r:(anios[r['anio']],longitudes[longitud(r)],actores[r['actor']],
                    reuniones[r['meeting_id']],rango(r['intervencion_id'],'seleccionar')))
            if any(casi_copia(firmas[r['intervencion_id']], f) for f in firmas_vetadas):
                rechazados_copia.add(r['intervencion_id']);progreso=True;continue
            firmas_vetadas.append(firmas[r['intervencion_id']])
            elegidos.append(r);usados.add(r['intervencion_id']);reuniones[r['meeting_id']]+=1
            actores[r['actor']]+=1;anios[r['anio']]+=1;longitudes[longitud(r)]+=1;progreso=True
        if not progreso: raise ValueError('Selección sin progreso')
    elegidos.sort(key=lambda r:rango(r['intervencion_id'],'presentacion'))
    for i,r in enumerate(elegidos,1):r['caso_id']=f'C{i:02d}';r['bloque']=(i-1)//10+1
    return elegidos, {'descartes':dict(descartes),'pool_con_pistas':len(pool),
        'casi_copias_descartadas':len(rechazados_copia),
        'pool_por_canal_tipo':{c+'_'+t:conteos[(c,t)] for c,t in celdas},
        'cuotas_efectivas':{c+'_'+t:cuotas[(c,t)] for c,t in celdas},
        'por_anio':dict(sorted(anios.items())),'por_longitud':dict(longitudes),
        'reuniones':len(reuniones),'actores':len(actores),'max_por_reunion':max(reuniones.values()),
        'max_por_actor':max(actores.values()),'caracteres':sum(len(r['texto']) for r in elegidos),
        'ids_anotados_excluidos':len(anotados),'ids_marco_humano_excluidos':len(humanos),
        'solapes_ids':len(usados & (anotados|humanos)),
        'solapes_textos_normalizados':sum(clave(r['texto']) in vetados for r in elegidos)}


def plan_folds(elegidos, corpus, folds):
    docs = {r['intervencion_id']:r for r in corpus}
    planes = {}
    for r in elegidos:
        prohibidos = [f['fold'] for f in folds if r['meeting_id'] in {docs[k]['meeting_id'] for k in f['validacion']}
                      or clave(r['texto']) in {clave(docs[k]['texto']) for k in f['validacion']}]
        planes[r['intervencion_id']] = {'folds_prohibidos':prohibidos,
            'folds_potencialmente_permitidos':[f['fold'] for f in folds if f['fold'] not in prohibidos],
            'incorporado_a_train':False}
    return planes


# ---- 3. Libro público: seis bloques, textos completos, sin sugerencias ocultas ----
def crear_excel(casos, guia, ruta):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.workbook.properties import CalcProperties
    X = cargar_script('29_preparar_excel_revision.py')
    exigir_salidas_nuevas(ruta)
    wb=Workbook();inicio=wb.active;inicio.title='Inicio'
    wb.properties.creator='Proyecto D&H';wb.properties.title='60 candidatos reales para revisión'
    wb.properties.created=wb.properties.modified=datetime(2026,9,16)
    wb.calculation=CalcProperties(calcMode='auto',fullCalcOnLoad=True)
    inicio.column_dimensions['A'].width=26;inicio.column_dimensions['B'].width=110
    X.titulo(inicio,1,'60 intervenciones reales · revisión por bloques','B')
    instrucciones=[
        'Guarda una copia antes de empezar y usa Ctrl+S / Cmd+S. Puedes devolver avances parciales.',
        'Empieza por Bloque 1: son diez casos. Completa únicamente las celdas amarillas.',
        'Pulsa Leer texto y lee TODAS las partes. El botón del caso permite volver a su respuesta.',
        'Elige Hawkish, Dovish o Neutral según la intervención completa. No intentes completar cuotas.',
        'No puedo decidir registra una duda, no una clase para entrenar. Explica el motivo.',
        'Relevante: Sí salvo formalidad/logística o texto no evaluable. Para No, usa Neutral y explica.',
        'Para H/D/N relevantes, pega una cita literal de hasta 300 caracteres. No copies explicaciones de otra IA.',
        'La muestra fue buscada con pistas de lenguaje. No contiene sugerencias por caso ni etiquetas ocultas.',
        'Estado comprueba campos, no si tu criterio o tu cita son correctos. Ninguna respuesta se incorpora sola al modelo.',
        'Consulta la Guía. Si recibes ayuda o viste sugerencias, decláralo; no reconstruyas respuestas perdidas.']
    for fila,t in enumerate(instrucciones,3):
        inicio.merge_cells(start_row=fila,start_column=1,end_row=fila,end_column=2)
        X.texto(inicio.cell(fila,1),t);inicio.row_dimensions[fila].height=36
    for b in range(1,7):
        X.enlace(inicio.cell(13+b,1),f'Bloque {b}: diez casos',f'Bloque {b}','A1')
    X.enlace(inicio['B14'],'Leer guía de decisión','Guía','A1')
    for fila,t in [(22,'Nombre (opcional)'),(23,'Fecha real de revisión'),(24,'Ayuda recibida'),(25,'¿Viste sugerencias previas?')]:
        X.texto(inicio.cell(fila,1),t);X.editable(inicio.cell(fila,2));inicio.row_dimensions[fila].height=30
    X.lista(inicio,'B24',['Ninguna','IA','Otra persona','IA y persona','Prefiero no declarar'])
    X.lista(inicio,'B25',['No','Sí','No recuerdo'])
    X.texto(inicio['A27'],'Identidad de la muestra');X.texto(inicio['B27'],VERSION)
    X.texto(inicio['A28'],'Huella de los textos')
    X.texto(inicio['B28'],hashlib.sha256(json.dumps([(r['caso_id'],r['texto']) for r in casos],ensure_ascii=False).encode()).hexdigest())
    ubicaciones={}
    for b in range(1,7):
        respuestas=wb.create_sheet(f'Bloque {b}');textos=wb.create_sheet(f'Textos {b}')
        X.titulo(respuestas,1,f'Bloque {b} · diez decisiones, sin sugerencias')
        X.enlace(respuestas['A2'],'Inicio','Inicio','A1');X.enlace(respuestas['H2'],'Guía','Guía','A1')
        headers=['Caso','Actor','Mi postura','Relevante','Cita literal (máx. 300)','Motivo / duda','Estado','Texto completo']
        for col,(t,ancho) in enumerate(zip(headers,[10,27,23,13,52,42,23,20]),1):
            c=X.texto(respuestas.cell(4,col),t);c.fill=PatternFill('solid',fgColor=X.AZUL)
            c.font=Font(name='Calibri',bold=True,color='FFFFFF');respuestas.column_dimensions[c.column_letter].width=ancho
        respuestas.freeze_panes='C5';respuestas.auto_filter.ref='A4:H14';respuestas.row_dimensions[4].height=32
        X.lista(respuestas,'C5:C14',['Hawkish','Dovish','Neutral','No puedo decidir'])
        X.lista(respuestas,'D5:D14',['Sí','No'])
        from openpyxl.worksheet.datavalidation import DataValidation
        limite=DataValidation(type='textLength',operator='lessThanOrEqual',formula1=300,allow_blank=True)
        limite.showErrorMessage=True;limite.errorStyle='stop';limite.error='Máximo 300 caracteres.'
        respuestas.add_data_validation(limite);limite.add('E5:E14')
        X.titulo(textos,1,f'Textos completos · bloque {b}','C')
        for col,ancho in [('A',12),('B',12),('C',112)]:textos.column_dimensions[col].width=ancho
        textos.freeze_panes='C3';fila_texto=3
        for fila,r in enumerate([r for r in casos if r['bloque']==b],5):
            X.texto(respuestas.cell(fila,1),r['caso_id']);X.texto(respuestas.cell(fila,2),r['actor'])
            for c in range(3,7):X.editable(respuestas.cell(fila,c))
            respuestas.cell(fila,7,X.formula_estado(fila));respuestas.row_dimensions[fila].height=98
            X.enlace(respuestas.cell(fila,8),'Leer '+r['caso_id'],textos.title,f'A{fila_texto}')
            X.enlace(textos.cell(fila_texto,1),'Volver '+r['caso_id'],respuestas.title,f'A{fila}')
            X.texto(textos.cell(fila_texto,3),' · '.join([r['actor'],r['cargo'],r['fecha']]))
            textos.row_dimensions[fila_texto].height=48;filas=[]
            for i,parte in enumerate(X.bloques(r['texto']),1):
                fila_texto+=1;filas.append(fila_texto)
                X.texto(textos.cell(fila_texto,1),r['caso_id']);X.texto(textos.cell(fila_texto,2),str(i))
                X.texto(textos.cell(fila_texto,3),parte)
                lineas=sum(max(1,math.ceil(len(l)/90)) for l in parte.split('\n'))
                altura=max(28,lineas*17+10)
                if altura>409:raise ValueError('Texto no visible en altura Excel')
                textos.row_dimensions[fila_texto].height=altura
            ubicaciones[r['caso_id']]={'hoja_respuesta':respuestas.title,'fila_respuesta':fila,
                                      'hoja_texto':textos.title,'filas_texto':filas}
            fila_texto+=2
    g=wb.create_sheet('Guía');g.column_dimensions['A'].width=4;g.column_dimensions['B'].width=120
    X.titulo(g,1,'Guía vigente y aclaraciones · leer el texto completo','B')
    X.enlace(g['B2'],'Inicio','Inicio','A1')
    for i,t in enumerate(guia.splitlines(),4):
        X.texto(g.cell(i,2),t);g.row_dimensions[i].height=max(22,math.ceil(max(1,len(t))/95)*17+8)
    for hoja in wb:
        X.proteger(hoja);hoja.sheet_properties.pageSetUpPr.fitToPage=True
        hoja.page_setup.orientation='landscape';hoja.page_setup.fitToWidth=1;hoja.page_setup.fitToHeight=0
    wb.save(ruta)
    return ubicaciones


# ---- 4. Creación exclusiva; evidencia técnica separada del libro del investigador ----
def ejecutar(salida=SALIDA):
    salida=Path(salida);exigir_salidas_nuevas(salida)
    corpus,anotados,humanos,folds,fuentes=cargar()
    casos,resumen=seleccionar(corpus,anotados,humanos)
    planes=plan_folds(casos,corpus,folds)
    guia=(RAIZ/'docs/codebook_v2.md').read_text(encoding='utf-8').split('## 4. Ejemplos semilla',1)[0]
    guia+='\nAclaraciones aceptadas: importa la dirección respaldada en toda la intervención, incluida la futura.\nMantener no implica automáticamente N; menor ritmo de alzas no implica D.\nNormalización no implica H por sí sola: identificar dirección y objeto.\nMención, historia, escenario condicional y alternativa rechazada no equivalen a adhesión.\nNo reparar OCR ni tomar tasas o dirección de otra intervención.\n'
    salida.mkdir(parents=True)
    ubicaciones=crear_excel(casos,guia,salida/NOMBRE)
    seleccion=[{'caso_id':r['caso_id'],'bloque':r['bloque'],'intervencion_id':r['intervencion_id'],
        'meeting_id':r['meeting_id'],'sha256_texto':hashlib.sha256(r['texto'].encode()).hexdigest(),
        'n_caracteres':len(r['texto']), 'canal_busqueda':r['canal_busqueda'],'tipo_busqueda':r['tipo_busqueda'],
        'coincidencias':r['coincidencias'],**planes[r['intervencion_id']],**ubicaciones[r['caso_id']]} for r in casos]
    def guardar(n,d):
        with (salida/n).open('x',encoding='utf-8') as f:json.dump(d,f,ensure_ascii=False,indent=2,allow_nan=False);f.write('\n')
    guardar('NO_CONSULTAR_antes_de_responder_seleccion.json',seleccion)
    guardar('resumen.json',{'version':VERSION,'semilla':SEMILLA,'n_candidatos':60,**resumen,
        'etiquetas_confirmadas':0,'entrenamiento_realizado':False,'datos_sinteticos':False,
        'deduplicacion':'IDs, texto normalizado y Jaccard de trigramas >=0.85 frente a excluidos y seleccionados; no equivalencia semántica completa',
        'muestra_representativa_o_test_independiente':False,'plan_folds_no_ejecutado':True})
    for p in [Path(__file__),RAIZ/'scripts/29_preparar_excel_revision.py',RAIZ/'scripts/40_gestionar_proyecto.py',
              RAIZ/'docs/codebook_v2.md',RAIZ/'docs/AMPLIACION_HD_60_V1.md']:
        fuentes[str(p.relative_to(RAIZ))]=sha256(p)
    guardar('manifest.json',{'version':VERSION,'sha256_fuentes':fuentes,
        'sha256_salidas':{p.name:sha256(p) for p in salida.iterdir() if p.is_file()}})
    print(json.dumps(resumen,ensure_ascii=False,indent=2));print('Excel preparado:',salida/NOMBRE)
    return casos


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--salida',type=Path,default=SALIDA)
    ejecutar(p.parse_args().salida)
