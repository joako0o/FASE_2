"""Compara la devolución humana de 30 casos con sus etiquetas IA de entrenamiento.

Acepta la hoja simplificada recibida (ID_30/texto/decisión/cita/motivo), conserva
la fuente y no importa correcciones. No abre las 306 respuestas del examen.
"""
# ---- 1. Entradas versionadas y normalización explícita, sin imputar decisiones ----
import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from utilidades import cargar_script, exigir_salidas_nuevas, norm, sha256

INTERFAZ = cargar_script('28_reparar_exportacion_revision.py')
RAIZ = Path(__file__).resolve().parents[1]
MUESTRA = INTERFAZ.SALIDA
DEVOLUCION = MUESTRA/'devolucion_humana_v1'
ORIGINAL = DEVOLUCION/'30_anotaciones_humanas.xlsx'
SALIDA = DEVOLUCION/'comparacion'
INFORME = RAIZ/'docs/REVISION_HUMANA_ENTRENAMIENTO_30_V1.md'
CLASES = ('hawkish','dovish','neutral')
ENCABEZADOS = ['ID_30','Texto','Tu decisión','Frase que respalda tu decisión','Motivo o duda']
DECISIONES = {**{c:c for c in CLASES},'no puedo decidir':'no_puedo_decidir'}
VACIOS = {'','-','–','—','.','...','…'}


def escribir_json(ruta,datos):
    with ruta.open('x',encoding='utf-8') as f:
        json.dump(datos,f,ensure_ascii=False,indent=2,allow_nan=False)
        f.write('\n')


def identificar(valor):
    if isinstance(valor,bool):
        raise ValueError('ID booleano')
    if isinstance(valor,(int,float)) and int(valor)==valor:
        numero=int(valor)
    elif isinstance(valor,str) and valor.strip().isdigit():
        numero=int(valor.strip())
    elif isinstance(valor,str) and valor.strip().startswith('R') and valor.strip()[1:].isdigit():
        numero=int(valor.strip()[1:])
    else:
        raise ValueError(f'ID no reconocido: {valor!r}')
    assert 1<=numero<=30,'ID fuera de muestra'
    return f'R{numero:02d}'


def calidad_cita(cita,texto):
    normalizada=norm(cita)
    if normalizada in VACIOS:
        return 'ausente_o_marcador'
    if normalizada not in norm(texto):
        return 'no_literal_y_larga' if len(cita)>300 else 'no_literal'
    return 'literal_larga' if len(cita)>300 else 'literal_hasta_300'


# ---- 2. Primero identidad por ID y texto; después se permite unir la clave IA ----
def leer_devolucion(ruta,publico):
    esperados={c['caso_id']:c for c in publico['casos']}
    assert len(esperados)==30
    libro=load_workbook(ruta,data_only=False,read_only=True,keep_links=False)
    try:
        assert len(libro.worksheets)==1,'Esta versión admite la hoja simplificada recibida'
        hoja=libro.worksheets[0]
        assert hoja.max_column==5
        filas=list(hoja.iter_rows())
        assert [c.value for c in filas[0]]==ENCABEZADOS,'Encabezados distintos'
        registros=[]
        for numero,fila in enumerate(filas[1:],2):
            if all(c.value is None for c in fila):
                continue
            assert all(c.data_type!='f' for c in fila),f'Fórmula en datos de fila {numero}'
            identificador=identificar(fila[0].value)
            assert identificador in esperados
            original=esperados[identificador]['texto']
            texto_recibido=fila[1].value
            assert isinstance(texto_recibido,str) and norm(texto_recibido)==norm(original),f'Texto no corresponde íntegramente: {identificador}'
            decision='' if fila[2].value is None else str(fila[2].value)
            normalizada=norm(decision).casefold()
            assert normalizada in DECISIONES or not normalizada,f'Decisión no reconocida: {identificador}'
            cita='' if fila[3].value is None else str(fila[3].value)
            nota='' if fila[4].value is None else str(fila[4].value)
            registros.append({'caso_id':identificador,'fila_excel':numero,'decision_original':decision,
                'etiqueta_humana':DECISIONES.get(normalizada,'sin_respuesta'),
                'relevancia_humana':'no_declarada','cita_humana':cita,'motivo_humano':nota,
                'texto_identico':texto_recibido==original,'texto_igual_tras_espacios':True,
                'sha256_texto_recibido':hashlib.sha256(texto_recibido.encode()).hexdigest(),
                'calidad_cita':calidad_cita(cita,original),'n_caracteres_cita':len(cita),
                'texto_original':original})
        tabla=pd.DataFrame(registros)
        assert len(tabla)==30 and tabla.caso_id.is_unique,'Faltan IDs o hay duplicados'
        assert set(tabla.caso_id)==set(esperados)
        return tabla.sort_values('caso_id').reset_index(drop=True)
    finally:
        libro.close()


def comparar(humanas,clave):
    assert set(clave.columns)=={'caso_id','intervencion_id','etiqueta_ia','es_relevante_ia','sha256_texto_original'}
    assert clave.caso_id.is_unique and set(clave.caso_id)==set(humanas.caso_id)
    assert clave.etiqueta_ia.isin(CLASES).all()
    tabla=humanas.merge(clave,on='caso_id',validate='one_to_one').sort_values('caso_id').reset_index(drop=True)
    for r in tabla.itertuples():
        assert hashlib.sha256(r.texto_original.encode()).hexdigest()==r.sha256_texto_original,r.caso_id
    definida=tabla.etiqueta_humana.isin(CLASES)
    tabla['acuerdo_postura']=tabla.etiqueta_humana.eq(tabla.etiqueta_ia) & definida
    tabla['comparacion_disponible']=definida
    sub=tabla[definida]
    matriz=[[int((sub.etiqueta_ia.eq(ia)&sub.etiqueta_humana.eq(h)).sum()) for h in CLASES] for ia in CLASES]
    n=len(sub);acuerdos=int(sub.acuerdo_postura.sum())
    resumen={'n_recibidas':len(tabla),'n_decisiones_comparables':n,'n_sin_decision_comparable':len(tabla)-n,
        'acuerdos':acuerdos,'desacuerdos':n-acuerdos,'proporcion_acuerdo':acuerdos/n if n else None,
        'distribucion_ia':tabla.etiqueta_ia.value_counts().to_dict(),
        'distribucion_humana':tabla.etiqueta_humana.value_counts().to_dict(),
        'matriz_filas_ia_columnas_humano':matriz,'orden_clases':list(CLASES),
        'acuerdo_por_estrato_ia':{c:{'n':int(tabla.etiqueta_ia.eq(c).sum()),'comparables':int(sub.etiqueta_ia.eq(c).sum()),
            'acuerdos':int((sub.etiqueta_ia.eq(c)&sub.acuerdo_postura).sum())} for c in CLASES},
        'calidad_citas':tabla.calidad_cita.value_counts().to_dict(),
        'textos_identicos':int(tabla.texto_identico.sum()),'textos_iguales_solo_tras_espacios':int((~tabla.texto_identico).sum()),
        'relevancia_humana':'No declarada: no se imputa con el flag IA.',
        'procedencia_revision':{'fecha_anotacion':'no_declarada','ayuda_recibida':'no_declarada','vio_etiquetas_ia':'no_declarado'},
        'importacion_canonica_realizada':False,'entrenamiento_realizado':False,'examen_306_abierto':False,
        'interpretacion':'Acuerdo sobre muestra balanceada por IA de entrenamiento; no accuracy de un modelo ni estimación poblacional.'}
    return tabla,resumen


# ---- 3. Informe de acuerdos, sin adjudicar automáticamente discrepancias ----
def informe(tabla,resumen):
    diferencias=tabla[tabla.comparacion_disponible & ~tabla.acuerdo_postura]
    lineas=['# Revisión humana del entrenamiento: 30 casos', '',
        f'**{resumen["acuerdos"]}/{resumen["n_decisiones_comparables"]} decisiones coinciden con la etiqueta IA; {resumen["desacuerdos"]} difieren.**', '',
        '**Esto compara anotaciones, no predicciones de un clasificador.** No es un nuevo test humano ni permite extrapolar una tasa de errores al corpus.', '',
        '## Recepción y vinculación', '',
        '- Original recibido desde el enlace GitHub proporcionado por el investigador, fijado al commit de `recepcion.json`, con Git blob SHA1 y SHA256 verificados. Se conserva sin cambios en `data/auditoria/revision_entrenamiento_30_v1/devolucion_humana_v1/`.',
        '- La devolución usa una hoja simplificada de cinco columnas: ID_30, texto, decisión, cita y motivo. No es el libro de cuatro hojas, pero la identidad se verifica por ID y texto completo, no por el orden supuesto de las filas.',
        f'- {resumen["textos_identicos"]} textos idénticos y {resumen["textos_iguales_solo_tras_espacios"]} iguales tras normalizar únicamente espacios. No se aceptan recortes, cambios léxicos o casos sustituidos.',
        '- No trae relevancia, fecha de anotación, ayuda recibida o declaración de acceso previo a etiquetas IA. Quedan **no declaradas**. El recibo UTC no se usa como fecha de anotación. No se afirma revisión ciega independiente.', '',
        '## Resultado de la comparación de postura', '',
        '| Etiqueta IA de origen | Humano H | Humano D | Humano N |',
        '|---|---:|---:|---:|']
    for clase,fila in zip(CLASES,resumen['matriz_filas_ia_columnas_humano']):
        lineas.append(f'| {clase} | {fila[0]} | {fila[1]} | {fila[2]} |')
    lineas+=['', 'Muestra seleccionada con 10 casos por clase **IA**, no humana. Una proporción global de acuerdo aquí está ponderada por ese diseño, no por la distribución real del corpus.', '',
        '| Estrato IA | Coincidencias / decisiones comparables |', '|---|---:|']
    for c in CLASES:
        v=resumen['acuerdo_por_estrato_ia'][c]
        lineas.append(f'| {c} | {v["acuerdos"]}/{v["comparables"]} |')
    lineas+=['', '## Casos que requieren revisión conjunta', '',
        'Se conserva tu decisión tal como la enviaste. Un desacuerdo no demuestra automáticamente que tú o la IA estén equivocados. Las citas se muestran sin sustituirlas por una elección del agente.', '',
        '| Caso | ID original | IA | Humano | Motivo del investigador |', '|---|---|---|---|---|']
    def celda(s):
        return str(s).replace('|','\\|').replace('\n',' ')
    for r in diferencias.itertuples():
        lineas.append(f'| {r.caso_id} | {r.intervencion_id} | {r.etiqueta_ia} | {r.etiqueta_humana} | {celda(r.motivo_humano)} |')
    if diferencias.empty:
        lineas+=['No hay discrepancias entre decisiones comparables.']
    lineas+=['', '## Citas y alcance documental', '',
        '| Control de cita | Casos |', '|---|---:|']
    for estado,n in resumen['calidad_citas'].items():
        lineas.append(f'| {estado} | {n} |')
    lineas+=['', 'Los marcadores «-» o «.» se registran como ausencia de cita, no como evidencia literal válida. La normalización para literalidad solo ajusta espacios; no corrige palabras o puntuación. Una cita literal tampoco garantiza que fundamente la postura.',
        'Estas incidencias **no invalidan ni cambian las decisiones de postura para esta comparación**. Sí impiden declarar completa una importación canónica bajo la guía v2, junto con la relevancia no declarada. No se rellenan los campos con las etiquetas/citas IA.', '',
        '## Próximo paso y límites', '',
        'Revisar las discrepancias con el texto completo y R1–R5: opción respaldada frente a menú, diagnóstico frente a postura, escenario extranjero y mantenimiento frente a sesgo. Documentar una decisión conjunta antes de proponer cualquier corrección, en una versión separada.',
        'No se entrenó BETO, no se modificaron las 1.352 etiquetas IA y no se abrió el examen de 306. La muestra procede de desarrollo, fue balanceada por IA y excluye los IDs/copias del examen y de la validación reutilizada según el protocolo original. Es una auditoría acotada, no una demostración de calidad global.', '',
        '## Reproducción', '',
        '`scripts/30_comparar_revision_humana.py` valida primero identidad y después une la clave IA. Rechaza sobrescribir. Guarda decisiones originales, citas y texto completo, incidencias, matriz y hashes de fuentes/salidas.', '',
        '```bash', 'python scripts/30_comparar_revision_humana.py --salida /ruta/nueva --informe /ruta/informe_nuevo.md', '```',
        'Los detalles de recepción, comparación y posterior repetición se conservan junto al original. La revisión cualitativa del agente, si se realiza, se documenta aparte y no sobrescribe las decisiones.', '']
    return '\n'.join(lineas)


def ejecutar(salida=SALIDA,ruta_informe=INFORME):
    salida,ruta_informe=Path(salida),Path(ruta_informe)
    exigir_salidas_nuevas(salida,ruta_informe)
    recibo=json.loads((DEVOLUCION/'recepcion.json').read_text())
    assert sha256(ORIGINAL)==recibo['sha256_original'],'Devolución original alterada'
    registro=json.loads((MUESTRA/'protocolo.json').read_text())
    archivo_publico=MUESTRA/'archivo_interfaz_v1/index.html'
    publico=INTERFAZ.payload(archivo_publico.read_text())
    assert publico['muestra_sha256']==registro['muestra_sha256']
    manifest=json.loads((MUESTRA/'archivo_interfaz_v1/manifest.json').read_text())
    assert sha256(archivo_publico)==manifest['sha256_salidas']['formulario/index.html']
    humanas=leer_devolucion(ORIGINAL,publico)
    ruta_clave=MUESTRA/'NO_ABRIR_hasta_finalizar_clave.csv'
    assert sha256(ruta_clave)==manifest['sha256_salidas'][ruta_clave.name]
    clave=pd.read_csv(ruta_clave,dtype=str,keep_default_na=False)
    tabla,resumen=comparar(humanas,clave)
    fuentes=[ORIGINAL,DEVOLUCION/'recepcion.json',MUESTRA/'protocolo.json',archivo_publico,ruta_clave,
        MUESTRA/'archivo_interfaz_v1/manifest.json',Path(__file__),RAIZ/'scripts/utilidades.py',
        RAIZ/'scripts/28_reparar_exportacion_revision.py',RAIZ/'tests/test_revision_humana_30.py',
        RAIZ/'docs/codebook_v2.md',RAIZ/'docs/PROTOCOLO_REVISION_ENTRENAMIENTO_30_V1.md']
    hashes={str(p.relative_to(RAIZ)):sha256(p) for p in fuentes}
    salida.mkdir(parents=True)
    tabla.to_csv(salida/'comparacion.csv',index=False)
    tabla[tabla.comparacion_disponible & ~tabla.acuerdo_postura].to_csv(salida/'desacuerdos.csv',index=False)
    tabla[tabla.calidad_cita.ne('literal_hasta_300')][['caso_id','calidad_cita','n_caracteres_cita','cita_humana']].to_csv(salida/'incidencias_citas.csv',index=False)
    escribir_json(salida/'resumen.json',resumen)
    ruta_informe.parent.mkdir(parents=True,exist_ok=True)
    with ruta_informe.open('x',encoding='utf-8') as f: f.write(informe(tabla,resumen))
    for nombre,h in hashes.items(): assert sha256(RAIZ/nombre)==h,nombre
    escribir_json(salida/'manifest.json',{'fecha_utc':datetime.now(timezone.utc).isoformat(),
        'sha256_insumos':hashes,'sha256_salidas':{p.name:sha256(p) for p in sorted(salida.iterdir()) if p.is_file()},
        'sha256_informe':sha256(ruta_informe),'original_humano_intacto':True,'entrenamiento_o_importacion':False})
    print(json.dumps(resumen,ensure_ascii=False,indent=2))
    return resumen


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--salida',type=Path,default=SALIDA)
    p.add_argument('--informe',type=Path,default=INFORME)
    args=p.parse_args()
    ejecutar(args.salida,args.informe)


if __name__=='__main__': main()
