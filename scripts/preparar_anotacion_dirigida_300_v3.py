"""Prepara 300 casos 2005–2015 dirigidos a D/H/N difícil, sin mostrar el estrato."""
import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data/muestras/anotacion_dirigida_300_v3_v1"
SEED = "dirigida-d150-h50-n100-v3-20260917"
QUOTAS = {"candidato_dovish": 150, "candidato_hawkish": 50, "neutral_dificil": 100}
MONETARY = re.compile(r"\b(tpm|tasa de pol[ií]tica monetaria|pol[ií]tica monetaria|tasa de instancia|impulso monetario|flap|facilidad de liquidez)\b", re.I)
DOVISH = re.compile(r"(sesgo (?:a|hacia) la baja|(?:bajar|reducir|recortar|disminuir|relajar|flexibilizar)\w*.{0,100}(?:tpm|tasa|pol[ií]tica)|(?:tpm|tasa|pol[ií]tica).{0,100}(?:bajar|reducir|recortar|disminuir|relajar|flexibilizar)\w*|pol[ií]tica monetaria expansiva|mayor impulso monetario|mantener\w*.{0,80}impulso monetario)", re.I)
HAWKISH = re.compile(r"(sesgo (?:a|hacia) el alza|(?:subir|elevar|aumentar|incrementar)\w*.{0,100}(?:tpm|tasa|pol[ií]tica)|(?:tpm|tasa|pol[ií]tica).{0,100}(?:subir|elevar|aumentar|incrementar)\w*|(?:retirar|reducir)\w*.{0,80}impulso monetario|normalizaci[oó]n de la pol[ií]tica monetaria|pol[ií]tica monetaria restrictiva)", re.I)
HARD_N = re.compile(r"\b(mantener|mantenci[oó]n|esperar|cautela|no descartar|depender[aá]|condicionado|expectativas? de mercado|compromiso|sin sesgo|riesgo inflacionario)\b", re.I)


def sha256(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def norm(text): return " ".join(text.split()).casefold()
def order(identity, salt=""): return hashlib.sha256(f"{SEED}|{salt}|{identity}".encode()).hexdigest()


def read_csv(path):
    with Path(path).open(encoding="utf-8", newline="") as f: return list(csv.DictReader(f))


def excluded_ids():
    paths = [ROOT / "data/evaluacion/referencia_v3/referencia_v3.csv",
             ROOT / "data/evaluacion/anotacion_adicional_300_v3_cerrada_v1/referencia_300_v3.csv"]
    ids = set()
    for path in paths: ids.update(r["intervencion_id"] for r in read_csv(path))
    # Evita volver a entregar instrumentos de anotación previos; no excluye tablas
    # administrativas como escalado_tandas, que contienen todo el universo.
    names = ["estrato_enriquecido.csv", "estrato_fases.csv", "estrato_tanda9.csv",
             "gold_ciego_300.csv", "piloto_300.csv", "test_retest_30.csv",
             "anotacion_adicional_300_v3/anotacion_adicional_300_v3.csv"]
    sample_paths = [ROOT / "data/muestras" / name for name in names]
    for path in sample_paths:
        rows = read_csv(path); ids.update(r["intervencion_id"] for r in rows)
    return ids, paths + sample_paths


def stratum(text):
    d, h, n, m = bool(DOVISH.search(text)), bool(HAWKISH.search(text)), bool(HARD_N.search(text)), bool(MONETARY.search(text))
    if d and not h: return "candidato_dovish"
    if h and not d: return "candidato_hawkish"
    if n and m and not d and not h: return "neutral_dificil"
    return ""


def select():
    corpus_path = ROOT / "data/L0/corpus.csv"; corpus = read_csv(corpus_path)
    excluded, sources = excluded_ids(); by_id = {r["intervencion_id"]: r for r in corpus}
    excluded_texts = {norm(by_id[i]["texto"]) for i in excluded if i in by_id}
    pools = {key: [] for key in QUOTAS}
    for row in corpus:
        if not (2005 <= int(row["anio"]) <= 2015) or row["flag_texto_danado"] == "True": continue
        if row["intervencion_id"] in excluded or norm(row["texto"]) in excluded_texts: continue
        key = stratum(row["texto"])
        if key: pools[key].append(row)
    for key in pools: pools[key].sort(key=lambda r: order(r["intervencion_id"], key))
    picked, meetings, years, selected_texts = [], Counter(), Counter(), set()
    # Cuotas se completan con topes suaves; si un estrato concentrado los agota, se relajan de forma fijada.
    for max_meeting, max_year in ((4, 38), (6, 45), (10, 60)):
        for key, quota in QUOTAS.items():
            have = sum(r["_estrato"] == key for r in picked)
            for row in pools[key]:
                if have >= quota: break
                text_key = norm(row["texto"])
                if row.get("_seleccionado") or text_key in selected_texts: continue
                if meetings[row["meeting_id"]] >= max_meeting or years[row["anio"]] >= max_year: continue
                item = dict(row); item["_estrato"] = key; row["_seleccionado"] = "1"
                picked.append(item); selected_texts.add(text_key); meetings[row["meeting_id"]] += 1; years[row["anio"]] += 1; have += 1
        if all(sum(r["_estrato"] == key for r in picked) == quota for key, quota in QUOTAS.items()): break
    counts = Counter(r["_estrato"] for r in picked)
    if counts != Counter(QUOTAS) or len({r["intervencion_id"] for r in picked}) != 300:
        raise ValueError(f"No se completaron cuotas: {counts}")
    picked.sort(key=lambda r: order(r["intervencion_id"], "blind-order"))
    return picked, corpus_path, sources, meetings, years


def make_workbook(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook(); start = wb.active; start.title = "Inicio"; sheet = wb.create_sheet("Anotación"); guide = wb.create_sheet("Codebook v3")
    navy, yellow, green = "173E4C", "FFF2CC", "DFF1E9"
    start["A1"] = "Anotación dirigida v3 · 300 intervenciones reales (2005–2015)"
    start["A1"].font = Font(size=16, bold=True, color="FFFFFF"); start["A1"].fill = PatternFill("solid", fgColor=navy)
    instructions = [
        "La selección buscó 150 candidatos dovish, 50 hawkish y 100 neutrales difíciles mediante vocabulario; esos estratos NO son etiquetas y están ocultos para evitar sesgo.",
        "No hay predicciones ni etiquetas sugeridas. No intentes reproducir cuotas: clasifica cada intervención completa según codebook v3.",
        "Completa etiqueta, relevancia, confianza, cita literal exacta (máximo 300 caracteres) y fundamento.",
        "Usa no_puedo_decidir cuando falte información indispensable y explica qué falta; no lo conviertas automáticamente en neutral.",
        "La muestra es ampliación de entrenamiento 2005–2015 y no un test final independiente."
    ]
    for i, text in enumerate(instructions, 3):
        start.cell(i, 1, text).alignment = Alignment(wrap_text=True, vertical="top"); start.row_dimensions[i].height = 34
    start.column_dimensions["A"].width = 120
    headers = ["orden", "intervencion_id", "fecha_reunion", "actor", "cargo", "texto", "etiqueta_v3", "es_relevante_v3", "confianza", "cita_literal", "fundamento", "estado"]
    widths = [8, 31, 15, 28, 26, 100, 19, 19, 14, 56, 56, 18]
    for col, (header, width) in enumerate(zip(headers, widths), 1):
        cell = sheet.cell(1, col, header); cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor=navy); cell.alignment = Alignment(wrap_text=True)
        sheet.column_dimensions[cell.column_letter].width = width
    for i, row in enumerate(rows, 2):
        values = [i - 1, row["intervencion_id"], row["fecha"], row["actor"], row["cargo"], row["texto"], "", "", "", "", "", "Pendiente"]
        for col, value in enumerate(values, 1):
            cell = sheet.cell(i, col, value); cell.alignment = Alignment(wrap_text=True, vertical="top")
            if 7 <= col <= 11: cell.fill = PatternFill("solid", fgColor=yellow); cell.protection = Protection(locked=False)
        sheet.cell(i, 12, f'=IF(OR(G{i}="",H{i}="",I{i}=""),"Pendiente",IF(G{i}="no_puedo_decidir",IF(K{i}="","Falta fundamento","Lista"),IF(H{i}="1",IF(OR(J{i}="",LEN(J{i})>300),"Revisar cita",IF(K{i}="","Falta fundamento","Lista")),IF(AND(H{i}="0",G{i}="neutral",K{i}<>""),"Lista","Revisar relevancia"))))')
        sheet.row_dimensions[i].height = 115
    for col, options in (("G", "hawkish,dovish,neutral,no_puedo_decidir"), ("H", "1,0"), ("I", "alta,media,baja")):
        validation = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True); sheet.add_data_validation(validation); validation.add(f"{col}2:{col}301")
    citation = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="300", allow_blank=True); sheet.add_data_validation(citation); citation.add("J2:J301")
    sheet.freeze_panes = "G2"; sheet.auto_filter.ref = "A1:L301"; sheet.protection.sheet = True; sheet.protection.selectUnlockedCells = True; sheet.protection.selectLockedCells = False
    codebook = (ROOT / "docs/codebook_v3.md").read_text(encoding="utf-8")
    guide.column_dimensions["A"].width = 120
    for i, line in enumerate(codebook.splitlines(), 1):
        guide.cell(i, 1, line).alignment = Alignment(wrap_text=True, vertical="top")
        if line.startswith("#"): guide.cell(i, 1).font = Font(bold=True, color=navy); guide.cell(i, 1).fill = PatternFill("solid", fgColor=green)
    guide.freeze_panes = "A2"; wb.active = 0; wb.save(path)


def write_csv(path, rows, fields):
    with path.open("x", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, lineterminator="\n"); writer.writeheader(); writer.writerows(rows)


def prepare(out=OUT):
    out = Path(out)
    if out.exists(): raise FileExistsError("No sobrescribir instrumento dirigido")
    rows, corpus, sources, meetings, years = select(); out.mkdir(parents=True)
    blind_fields = ["orden", "intervencion_id", "fecha_reunion", "actor", "cargo", "texto", "etiqueta_v3", "es_relevante_v3", "confianza", "cita_literal", "fundamento", "estado"]
    blind = [{"orden": i, "intervencion_id": r["intervencion_id"], "fecha_reunion": r["fecha"], "actor": r["actor"], "cargo": r["cargo"], "texto": r["texto"], "etiqueta_v3": "", "es_relevante_v3": "", "confianza": "", "cita_literal": "", "fundamento": "", "estado": "pendiente"} for i, r in enumerate(rows, 1)]
    control = [{"orden": i, "intervencion_id": r["intervencion_id"], "estrato_recuperacion_no_visible": r["_estrato"], "meeting_id": r["meeting_id"], "anio": r["anio"], "sha256_texto": hashlib.sha256(r["texto"].encode()).hexdigest()} for i, r in enumerate(rows, 1)]
    write_csv(out / "anotacion_dirigida_300_v3.csv", blind, blind_fields)
    write_csv(out / "control_muestreo.csv", control, list(control[0]))
    make_workbook(rows, out / "anotacion_dirigida_300_v3.xlsx")
    summary = {"filas": 300, "ids_unicos": 300, "periodo": [2005, 2015], "cuotas_recuperacion": QUOTAS,
        "advertencia": "cuotas léxicas de recuperación, no etiquetas", "estrato_visible_en_excel": False,
        "etiquetas_visibles": False, "predicciones_consultadas_o_visibles": False,
        "reuniones": len(meetings), "maximo_por_reunion": max(meetings.values()), "anios": dict(sorted(years.items())),
        "respuestas_pendientes": 300, "rol": "ampliacion de entrenamiento real 2005-2015"}
    (out / "resumen.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    protocol = {"operacion": "muestreo determinista ciego por expresiones prefijadas, sin etiquetas ni predicciones",
        "prioridad_excluyente": ["candidato_dovish", "candidato_hawkish", "neutral_dificil"], "seed": SEED,
        "exclusiones": "todas las referencias v3, las 300 humanas cerradas y todos los instrumentos CSV previos",
        "fuentes_sha256": {str(path.relative_to(ROOT)): sha256(path) for path in [corpus, ROOT / "docs/codebook_v3.md", Path(__file__), *sources]}}
    (out / "protocolo.json").write_text(json.dumps(protocol, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    output_names = ["anotacion_dirigida_300_v3.csv", "anotacion_dirigida_300_v3.xlsx", "control_muestreo.csv", "resumen.json", "protocolo.json"]
    (out / "manifest.json").write_text(json.dumps({"sha256_salidas": {name: sha256(out / name) for name in output_names}}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return summary


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__); parser.add_argument("--salida", type=Path, default=OUT)
    print(json.dumps(prepare(parser.parse_args().salida), ensure_ascii=False, indent=2))
