"""Cierra las 300 anotaciones ciegas tras validación humana y repara citas administrativas."""
import csv
import hashlib
import json
import re
import shutil
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data/recepcion/evaluacion_ciega_300_v3_v1/evaluacion_ciega_300_v3_recibida.xlsx"
PROTECTED = ROOT / "data/muestras/evaluacion_ciega_300_v3_v1/evaluacion_ciega_300_v3.csv"
OUT = ROOT / "data/evaluacion/evaluacion_ciega_300_v3_cerrada_v1"
SPECIAL_MARKERS = {
    47: "el conjunto de antecedentes",
    98: "el Consejo del Banco Central acordó aumentar",
    99: "el Consejo del Banco Central acordó aumentar",
    244: "PMI global",
    250: "el Consejo del Banco Central de Chile acordó aumentar",
    295: "El Consejo aprobó",
}
ADJUDICATIONS = {
    1: {"etiqueta_v3": "N", "es_relevante_v3": "1", "confianza": "media", "fundamento": "Advertencia metodológica sustantiva, sin acción, recomendación ni orientación monetaria propia; corresponde N."},
    32: {"etiqueta_v3": "H", "es_relevante_v3": "1", "confianza": "alta", "fundamento": "Aunque vota mantener, hace propio un sesgo explícito hacia el inicio muy cercano de la normalización. La orientación futura respaldada cuenta y corresponde H."},
    256: {"etiqueta_v3": "no_puedo_decidir", "es_relevante_v3": "1", "confianza": "baja", "fundamento": "El fragmento no identifica el objeto de la proyección; falta información indispensable y no debe forzarse N."},
}
LABEL_MAP = {"H": "hawkish", "D": "dovish", "N": "neutral", "no_puedo_decidir": "no_puedo_decidir"}


def norm(value): return re.sub(r"\s+", " ", str(value or "")).strip()
def sha(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def write_json(path, obj): path.write_text(json.dumps(obj, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def exact_citation(order, citation, text):
    citation, text = norm(citation), norm(text)
    if not citation: return ""
    if citation in text: return citation
    stripped = re.sub(r"\s*\.{3}$", "", citation).rstrip()
    if stripped in text: return stripped
    marker = SPECIAL_MARKERS.get(order)
    if not marker: raise ValueError(f"Cita no literal sin reparación congelada: orden {order}")
    start = text.find(marker)
    if start < 0: raise ValueError(f"Marcador ausente: orden {order}")
    candidate = text[start:start + 300]
    if len(candidate) == 300 and start + 300 < len(text): candidate = candidate.rsplit(" ", 1)[0]
    return candidate


def run(out=OUT):
    out = Path(out)
    if out.exists(): raise FileExistsError("No sobrescribir cierre")
    wb = load_workbook(SOURCE); ws = wb["Anotación"]
    headers = [c.value for c in ws[1]]; index = {name: i + 1 for i, name in enumerate(headers)}
    protected = {}
    with PROTECTED.open(encoding="utf-8", newline="") as f:
        for row in csv.DictReader(f): protected[row["intervencion_id"]] = row["texto"]
    records, repairs, changes = [], [], []
    for excel_row in range(2, ws.max_row + 1):
        row = {h: ws.cell(excel_row, col).value for h, col in index.items()}; order = int(row["orden"]); iid = str(row["intervencion_id"])
        if iid not in protected or norm(row["texto"]) != norm(protected[iid]): raise ValueError(f"Texto/ID alterado: {iid}")
        old_label = str(row["etiqueta_v3"] or "")
        if order in ADJUDICATIONS:
            for field, value in ADJUDICATIONS[order].items(): ws.cell(excel_row, index[field], value=value); row[field] = value
            changes.append({"orden": order, "intervencion_id": iid, "etiqueta_recibida": old_label, "etiqueta_final": row["etiqueta_v3"], "decision": "confirmada por usuario 2026-09-17"})
        fixed = exact_citation(order, row["cita_literal"], row["texto"])
        if fixed != norm(row["cita_literal"]):
            repairs.append({"orden": order, "intervencion_id": iid, "cita_recibida": norm(row["cita_literal"]), "cita_final": fixed})
            ws.cell(excel_row, index["cita_literal"], value=fixed); row["cita_literal"] = fixed
        final_label = str(row["etiqueta_v3"])
        included = final_label in {"H", "D", "N"}
        ws.cell(excel_row, index["estado"], value="cerrado" if included else "cerrado_no_decidible")
        citation = norm(row["cita_literal"])
        if str(row["es_relevante_v3"]) == "1" and (not citation or citation not in norm(row["texto"])): raise ValueError(f"Cita final inválida: {iid}")
        records.append({"orden": order, "intervencion_id": iid, "meeting_id": iid.rsplit(":", 2)[0], "fecha_reunion": str(row["fecha_reunion"]), "actor": str(row["actor"] or ""), "cargo": str(row["cargo"] or ""), "texto": str(row["texto"]), "etiqueta_v3": LABEL_MAP[final_label], "es_relevante_v3": str(row["es_relevante_v3"]), "confianza": str(row["confianza"]), "cita_literal": citation, "fundamento": str(row["fundamento"]), "estado_revision_v3": "cerrado" if included else "cerrado_no_decidible", "incluir_evaluacion": str(included).lower(), "sha256_texto": hashlib.sha256(str(row["texto"]).encode()).hexdigest(), "procedencia": "IA asistida; validación humana fila por fila"})
    if len(records) != 300 or len({r["intervencion_id"] for r in records}) != 300 or len(repairs) != 33: raise ValueError("Conteos finales inesperados")
    out.mkdir(parents=True); wb.save(out / "evaluacion_ciega_300_v3_cerrada.xlsx")
    for name, rows in [("referencia_ciega_300_v3.csv", records), ("reparaciones_citas.csv", repairs), ("adjudicaciones_finales.csv", changes)]:
        with (out / name).open("x", encoding="utf-8", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=list(rows[0]), lineterminator="\n"); writer.writeheader(); writer.writerows(rows)
    summary = {"filas_recibidas": 300, "filas_con_decision_humana": 300, "filas_evaluables": 299, "no_decidibles": 1, "distribucion_evaluable": dict(Counter(r["etiqueta_v3"] for r in records if r["incluir_evaluacion"] == "true")), "reparaciones_citas": len(repairs), "adjudicaciones_pendientes_confirmadas": 3, "procedencia": "IA asistida con validación humana completa confirmada por el usuario", "evaluacion_modelos_abierta": False}
    write_json(out / "resumen.json", summary)
    outputs = ["evaluacion_ciega_300_v3_cerrada.xlsx", "referencia_ciega_300_v3.csv", "reparaciones_citas.csv", "adjudicaciones_finales.csv", "resumen.json"]
    write_json(out / "manifest.json", {"fuentes_sha256": {str(SOURCE.relative_to(ROOT)): sha(SOURCE), str(PROTECTED.relative_to(ROOT)): sha(PROTECTED)}, "sha256_salidas": {name: sha(out / name) for name in outputs}})
    return summary


if __name__ == "__main__": print(json.dumps(run(), ensure_ascii=False, indent=2))
