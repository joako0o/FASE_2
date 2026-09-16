"""Valida e importa una devolución gold sin modificar fuentes ni forzar errores.

El marco canónico es inmutable. La salida conserva sus textos y metadatos,
incorpora las cinco respuestas y una fecha de anotación explícita. No es aún
el esquema L1 con probabilidades: para κ bastan clases emparejadas por ID.

Uso: python scripts/fusionar_gold_llenado.py archivo.xlsx --validar
     python scripts/fusionar_gold_llenado.py archivo_corregido.xlsx --salida nuevo.csv
Fechas confirmadas por el anotador: --fecha-anotacion AAAA-MM-DD (opcional).
Una entrega incompleta requiere --permitir-parcial. No existe --forzar.
"""

# ---- 1. Parámetros y lectura sin pérdida de saltos de línea ----
import argparse
import csv
from datetime import date, datetime, timezone
import io
import json
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

import config
from utilidades import errores_anotacion, exigir_salidas_nuevas, norm, sha256

CANONICO = config.RUTA_MUESTRAS / "gold_ciego_300.csv"
SALIDA_DEFECTO = config.RUTA_MUESTRAS / "gold_ciego_300_llenado.csv"
COLS_LLENAR = ["etiqueta", "confianza", "es_relevante", "nota", "frase_justificante"]
CONFIANZAS_GOLD = {"alta", "media"}


def texto_celda(valor):
    """Convierte tipos Excel sin inventar fechas ni aceptar fórmulas como valores."""
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.date().isoformat() if valor.time() == datetime.min.time() else str(valor)
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, bool):
        return str(int(valor))
    if isinstance(valor, (int, float)) and float(valor).is_integer():
        return str(int(valor))
    return str(valor)


def leer_xlsx(ruta):
    """Rechaza fórmulas: no usa caches que podrían estar vacíos o desactualizados."""
    with Path(ruta).open("rb") as archivo:
        libro = load_workbook(archivo, read_only=True, data_only=False)
        try:
            assert "etiquetar" in libro.sheetnames, "falta hoja etiquetar"
            filas = []
            for fila in libro["etiquetar"].iter_rows():
                assert not any(c.data_type == "f" for c in fila), "el instrumento no admite fórmulas"
                filas.append([texto_celda(c.value) for c in fila])
        finally:
            libro.close()
    assert filas, "archivo vacío"
    return [c.strip() for c in filas[0]], [f for f in filas[1:] if any(c.strip() for c in f)], "xlsx"


def leer_tabla(ruta):
    """CSV con BOM y coma/punto y coma/tabulador; conserva saltos internos."""
    ruta = Path(ruta)
    if ruta.suffix.lower() in {".xlsx", ".xlsm"}:
        return leer_xlsx(ruta)
    crudo = ruta.read_bytes()
    try:
        texto = crudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = crudo.decode("cp1252")
        print("AVISO: CSV leído como cp1252; verificar tildes.")
    for separador in [";", ",", "\t"]:
        lector = csv.reader(io.StringIO(texto, newline=""), delimiter=separador)
        cabecera = next(lector, [])
        if "intervencion_id" in cabecera:
            return cabecera, [f for f in lector if any(c.strip() for c in f)], separador
    raise ValueError("CSV sin cabecera reconocible")


def registros(ruta):
    """Controla cabeceras duplicadas y filas irregulares antes de mapear columnas."""
    cabecera, filas, _ = leer_tabla(ruta)
    assert cabecera and len(cabecera) == len(set(cabecera)), "cabecera vacía o duplicada"
    assert all(len(f) == len(cabecera) for f in filas), "filas con distinto número de columnas"
    return cabecera, [dict(zip(cabecera, f)) for f in filas]


# ---- 2. Validación estructural, de respuestas y fechas ----
def fecha_valida(valor, corte):
    """No acepta fechas futuras ni formatos ambiguos de Excel/localización."""
    try:
        fecha = date.fromisoformat(valor)
        return fecha.isoformat() == valor and fecha <= corte
    except ValueError:
        return False


def analizar(cabecera, canonico, cabecera_lleno, llenado, fecha_anotacion=None, corte=None):
    """Devuelve filas candidatas, incidencias y pendientes; nunca escribe ni corrige citas."""
    corte = corte or datetime.now(ZoneInfo("America/Santiago")).date()
    if fecha_anotacion is not None:
        assert fecha_valida(fecha_anotacion, corte), "fecha-anotacion inválida o futura"
    assert cabecera_lleno == cabecera, "columnas distintas del instrumento canónico"
    ids_c = [r["intervencion_id"] for r in canonico]
    ids_l = [r["intervencion_id"] for r in llenado]
    assert len(ids_c) == len(set(ids_c)), "IDs duplicados en canónico"
    assert len(ids_l) == len(set(ids_l)), "IDs duplicados en devolución"
    assert set(ids_c) == set(ids_l), "IDs distintos del canónico"
    por_id = {r["intervencion_id"]: r for r in llenado}
    nuevas, incidencias, pendientes = [], [], []
    for original in canonico:
        iid = original["intervencion_id"]
        lleno = por_id[iid]
        problemas = []
        for col in cabecera:
            if col not in COLS_LLENAR + ["fecha"] and norm(lleno[col]) != norm(original[col]):
                problemas.append(f"columna protegida modificada: {col}")
        valores = {c: lleno[c].strip() for c in COLS_LLENAR}
        # Normalización controlada que también se refleja en la salida.
        valores["etiqueta"] = valores["etiqueta"].lower()
        valores["confianza"] = valores["confianza"].lower()
        nueva = {**original, **valores}
        if not any(valores.values()):
            pendientes.append(iid)
            nueva["fecha"] = ""
        else:
            problemas += errores_anotacion(valores["etiqueta"], valores["confianza"],
                                            valores["es_relevante"], valores["nota"],
                                            valores["frase_justificante"], original["texto"],
                                            confianzas=CONFIANZAS_GOLD)
            nueva["fecha"] = fecha_anotacion if fecha_anotacion is not None else lleno["fecha"].strip()
            if not fecha_valida(nueva["fecha"], corte):
                problemas.append("fecha de anotación ausente, inválida o futura")
        for problema in problemas:
            incidencias.append({"orden": original["orden"], "intervencion_id": iid,
                                "problema": problema})
        nuevas.append(nueva)
    return nuevas, incidencias, pendientes


# ---- 3. Escritura explícita solo después de validar todo ----
def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("archivo", type=Path)
    parser.add_argument("--salida", type=Path, default=SALIDA_DEFECTO)
    parser.add_argument("--validar", action="store_true", help="solo informa, nunca escribe")
    parser.add_argument("--permitir-parcial", action="store_true")
    parser.add_argument("--fecha-anotacion", help="fecha real confirmada por el anotador para todas las respuestas")
    args = parser.parse_args(argv)
    try:
        cabecera, canonico = registros(CANONICO)
        cab_lleno, llenado = registros(args.archivo)
        nuevas, errores, pendientes = analizar(cabecera, canonico, cab_lleno, llenado,
                                               fecha_anotacion=args.fecha_anotacion)
        print(f"Respuestas: {len(nuevas) - len(pendientes)}/{len(nuevas)}; incidencias: {len(errores)}")
        for error in errores:
            print(f"  orden {error['orden']} ({error['intervencion_id']}): {error['problema']}")
        if errores or (pendientes and not args.permitir_parcial):
            print("No se escribió salida. Corregir incidencias; los pendientes requieren --permitir-parcial.")
            return 1
        if args.validar:
            print("Validación OK; modo solo lectura.")
            return 0
        manifiesto = args.salida.with_suffix(".provenance.json")
        exigir_salidas_nuevas(args.salida, manifiesto)
        assert args.salida.resolve() not in {CANONICO.resolve(), args.archivo.resolve()}, "salida sobre fuente"
        args.salida.parent.mkdir(parents=True, exist_ok=True)
        with args.salida.open("x", newline="", encoding="utf-8") as archivo:
            escritor = csv.DictWriter(archivo, fieldnames=cabecera, lineterminator="\n")
            escritor.writeheader()
            escritor.writerows(nuevas)
        trazabilidad = {"fecha_importacion_utc": datetime.now(timezone.utc).isoformat(),
                        "version_codebook": "v2", "entrada": str(args.archivo),
                        "sha256_entrada": sha256(args.archivo), "sha256_canonico": sha256(CANONICO),
                        "sha256_salida": sha256(args.salida), "filas": len(nuevas),
                        "pendientes": len(pendientes), "fecha_anotacion_confirmada": args.fecha_anotacion}
        with manifiesto.open("x", encoding="utf-8") as archivo:
            json.dump(trazabilidad, archivo, ensure_ascii=False, indent=2)
        print(f"OK: {args.salida}. Instrumento validado, no esquema L1 con probabilidades.")
        return 0
    except (AssertionError, ValueError, OSError) as error:
        print(f"ERROR: {error}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
