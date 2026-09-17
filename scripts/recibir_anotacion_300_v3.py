"""Recibe, normaliza y audita estructuralmente las 300 anotaciones humanas v3."""
import csv
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data/recepcion/anotacion_adicional_300_v3_v1"
SOURCE = OUT / "anotacion_original.xlsx"
BLIND = ROOT / "data/muestras/anotacion_adicional_300_v3/anotacion_adicional_300_v3.csv"
LABELS = {"H": "hawkish", "D": "dovish", "N": "neutral",
          "hawkish": "hawkish", "dovish": "dovish", "neutral": "neutral"}
DIRECTION_RE = re.compile(r"\b(subir|sube|suba|alza|alzas|elevar|aumentar|incrementar|endurec|restrictiv|"
                          r"bajar|baja|bajas|rebaj|reducir|reducci|recort|relaj|expansiv|est[ií]mulo|"
                          r"menos expansiv|m[aá]s expansiv)\w*", re.I)
MANUAL_MARKERS = {
    "RPM-2014-09-11:6395:1": "al Banco Central Europeo a adoptar nuevas bajas de tasas y a comprar activos.",
    "RPM-2009-07-09:2646:2": "concuerda también con los planteamientos acerca de la necesidad de impulsar hacia abajo",
}


def norm(value):
    return " ".join(str("" if value is None else value).split())


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def exact_case_insensitive(quote, text):
    start = text.lower().find(quote.lower())
    return text[start:start + len(quote)] if start >= 0 else None


def repair_quote(identity, quote, text):
    q, t = norm(quote), norm(text)
    if q in t:
        return q, "sin_cambio"
    trimmed = q[:-3].rstrip() if q.endswith("...") else q
    exact = exact_case_insensitive(trimmed, t)
    if exact:
        reason = "elipsis_no_literal_retirada" if q.endswith("...") else "capitalizacion_repuesta_desde_texto"
        return exact, reason
    marker = MANUAL_MARKERS.get(identity)
    if marker:
        start = t.find(marker)
        if start < 0:
            raise ValueError(f"Marcador manual ausente: {identity}")
        # Conservar una cita continua cercana a la longitud humana, nunca superior a 300.
        end = min(len(t), start + min(300, max(len(q), len(marker))))
        candidate = t[start:end].rstrip()
        return candidate, "transcripcion_repuesta_desde_texto"
    raise ValueError(f"Cita no reparable automáticamente: {identity}")


def run(out=OUT):
    out = Path(out)
    csv_path = out / "anotaciones_canonicas.csv"
    queue_path = out / "revision_semantica_prioritaria_100.csv"
    incidents_path = out / "incidencias_reparadas.csv"
    summary_path = out / "resumen.json"
    protocol_path = out / "protocolo.json"
    manifest_path = out / "manifest.json"
    generated = [csv_path, queue_path, incidents_path, summary_path, protocol_path, manifest_path]
    if any(p.exists() for p in generated):
        raise FileExistsError("No sobrescribir recepción existente")
    out.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(SOURCE, read_only=True, data_only=False)
    if wb.sheetnames != ["Inicio", "Anotación", "Codebook v3"]:
        raise ValueError("Hojas inesperadas")
    ws = wb["Anotación"]
    headers = [c.value for c in ws[1]]
    expected = ["orden", "intervencion_id", "fecha_reunion", "actor", "cargo", "texto",
                "etiqueta_v3", "es_relevante_v3", "confianza", "cita_literal", "fundamento", "estado"]
    if headers != expected or ws.max_row != 301:
        raise ValueError("Esquema o número de filas inesperado")
    human = [dict(zip(headers, values)) for values in ws.iter_rows(min_row=2, max_row=301, values_only=True)]
    with BLIND.open(encoding="utf-8", newline="") as f:
        blind = {row["intervencion_id"]: row for row in csv.DictReader(f)}
    if len(human) != 300 or len({r["intervencion_id"] for r in human}) != 300:
        raise ValueError("Se requieren 300 IDs únicos")

    canonical, incidents = [], []
    immutable = ["orden", "fecha_reunion", "actor", "cargo", "texto"]
    for row in human:
        identity = norm(row["intervencion_id"]); original = blind.get(identity)
        if original is None:
            raise ValueError(f"ID ajeno a muestra ciega: {identity}")
        for field in immutable:
            if norm(row[field]) != norm(original[field]):
                raise ValueError(f"Campo fuente modificado: {identity}/{field}")
        raw_label = norm(row["etiqueta_v3"])
        if raw_label not in LABELS:
            raise ValueError(f"Etiqueta inválida: {identity}/{raw_label}")
        label = LABELS[raw_label]
        relevance = norm(row["es_relevante_v3"])
        confidence = norm(row["confianza"])
        rationale = norm(row["fundamento"])
        if relevance not in {"0", "1"} or confidence not in {"alta", "media", "baja"} or not rationale:
            raise ValueError(f"Respuesta incompleta: {identity}")
        if relevance == "0" and label != "neutral":
            raise ValueError(f"Dirección en fila irrelevante: {identity}")
        raw_quote = norm(row["cita_literal"])
        if relevance == "1" and not raw_quote:
            raise ValueError(f"Falta cita relevante: {identity}")
        if relevance == "0":
            quote, repair = "", "sin_cambio"
        else:
            quote, repair = repair_quote(identity, raw_quote, row["texto"])
            if len(quote) > 300 or quote not in norm(row["texto"]):
                raise ValueError(f"Cita reparada inválida: {identity}")
            if repair != "sin_cambio":
                incidents.append({"intervencion_id": identity, "tipo": repair,
                                  "cita_original": raw_quote, "cita_reparada": quote})
        canonical.append({"orden": int(row["orden"]), "intervencion_id": identity,
            "fecha_reunion": norm(row["fecha_reunion"]), "actor": norm(row["actor"]),
            "cargo": norm(row["cargo"]), "texto": norm(row["texto"]),
            "etiqueta_humana_original": raw_label, "etiqueta_v3": label,
            "es_relevante_v3": relevance, "confianza_humana": confidence,
            "cita_literal": quote, "fundamento_humano": rationale,
            "estado_original": norm(row["estado"]), "estado_estructural": "valido",
            "estado_semantico": "pendiente_segunda_revision",
            "sha256_texto": hashlib.sha256(norm(row["texto"]).encode()).hexdigest()})

    with csv_path.open("x", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=canonical[0], lineterminator="\n")
        writer.writeheader(); writer.writerows(canonical)
    with incidents_path.open("x", encoding="utf-8", newline="") as f:
        fields = ["intervencion_id", "tipo", "cita_original", "cita_reparada"]
        writer = csv.DictWriter(f, fieldnames=fields, lineterminator="\n")
        writer.writeheader(); writer.writerows(incidents)

    # Segunda revisión: las 80 direccionales, más 20 N de mayor riesgo.
    directional = [r for r in canonical if r["etiqueta_v3"] != "neutral"]
    neutral_risk = [r for r in canonical if r["etiqueta_v3"] == "neutral" and
                    (r["confianza_humana"] == "baja" or DIRECTION_RE.search(r["texto"]))]
    neutral_risk.sort(key=lambda r: (r["confianza_humana"] != "baja", r["confianza_humana"] == "alta", r["orden"]))
    queue = directional + neutral_risk[:20]
    if len(directional) != 80 or len(queue) != 100:
        raise ValueError("Cola semántica esperada 80 direccionales + 20 N")
    queue_fields = ["intervencion_id", "fecha_reunion", "actor", "cargo", "texto", "etiqueta_v3",
                    "es_relevante_v3", "confianza_humana", "cita_literal", "fundamento_humano",
                    "revision_etiqueta_v3", "revision_relevancia_v3", "revision_confianza",
                    "revision_cita_literal", "revision_fundamento", "revision_estado"]
    with queue_path.open("x", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=queue_fields, lineterminator="\n")
        writer.writeheader()
        for r in queue:
            writer.writerow({**{k: r[k] for k in queue_fields if k in r},
                             **{k: "" for k in queue_fields if k.startswith("revision_")}})

    from collections import Counter
    summary = {"filas": 300, "ids_unicos": 300,
        "etiquetas_v3": dict(Counter(r["etiqueta_v3"] for r in canonical)),
        "relevancia_v3": dict(Counter(r["es_relevante_v3"] for r in canonical)),
        "confianza": dict(Counter(r["confianza_humana"] for r in canonical)),
        "citas_reparadas": len(incidents), "campos_fuente_modificados": 0,
        "filas_estructuralmente_validas": 300, "filas_semanticamente_cerradas": 0,
        "cola_segunda_revision": 100, "lista_para_entrenamiento": False}
    summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    protocol = {"origen_github": "main@8f1eedf:anotacion_adicional_300_v3.xlsx",
        "sha256_origen": sha(SOURCE), "sha256_muestra_ciega": sha(BLIND),
        "normalizacion_etiquetas": {"H": "hawkish", "D": "dovish", "N": "neutral"},
        "reparacion_citas": "solo extracción literal del mismo texto; sin cambiar etiqueta ni fundamento",
        "regla_revision": "100 casos: las 80 H/D y 20 N de riesgo; sin predicciones de modelo",
        "entrenamiento_permitido": False}
    protocol_path.write_text(json.dumps(protocol, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    manifest_path.write_text(json.dumps({"sha256_salidas": {p.name: sha(p) for p in
        [csv_path, queue_path, incidents_path, summary_path, protocol_path]}}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return summary


if __name__ == "__main__":
    print(json.dumps(run(), ensure_ascii=False, indent=2))
