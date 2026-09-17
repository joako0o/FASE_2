"""Prepara evaluación ciega 2005–2015 con reuniones completas en cuarentena."""
import argparse
import csv
import hashlib
import json
from collections import Counter
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

try:
    from preparar_anotacion_dirigida_300_v3 import stratum
except ModuleNotFoundError:  # importación como módulo desde tests
    from scripts.preparar_anotacion_dirigida_300_v3 import stratum

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data/muestras/evaluacion_ciega_300_v3_v1"
SEED = "evaluacion-ciega-grupos-v3-20260917"
REFERENCE_PATHS = [
    ROOT / "data/evaluacion/referencia_v3/referencia_v3.csv",
    ROOT / "data/evaluacion/anotacion_adicional_300_v3_cerrada_v1/referencia_300_v3.csv",
    ROOT / "data/evaluacion/anotacion_dirigida_300_v3_cerrada_v1/referencia_300_v3.csv",
]
INSTRUMENT_PATHS = [ROOT / "data/muestras" / name for name in ["estrato_enriquecido.csv", "estrato_fases.csv", "estrato_tanda9.csv", "gold_ciego_300.csv", "piloto_300.csv", "test_retest_30.csv"]]


def sha(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def rank(value, salt=""): return hashlib.sha256(f"{SEED}|{salt}|{value}".encode()).hexdigest()


def select():
    corpus_path = ROOT / "data/L0/corpus.csv"; corpus = pd.read_csv(corpus_path, dtype=str, keep_default_na=False)
    excluded = set()
    for path in REFERENCE_PATHS + INSTRUMENT_PATHS: excluded.update(pd.read_csv(path, dtype=str, keep_default_na=False).intervencion_id)
    remaining = corpus[~corpus.intervencion_id.isin(excluded) & ~corpus.flag_texto_danado.eq("True")].copy()
    remaining = remaining[remaining.anio.astype(int).between(2005, 2015)].copy(); remaining["estrato"] = remaining.texto.map(stratum)
    pivot = remaining.pivot_table(index=["anio", "meeting_id"], columns="estrato", values="intervencion_id", aggfunc="count", fill_value=0).reset_index()
    for column in ["candidato_dovish", "candidato_hawkish", "neutral_dificil"]:
        if column not in pivot: pivot[column] = 0
    pivot["score"] = pivot.candidato_dovish * 10 + pivot.candidato_hawkish * 2 + pivot.neutral_dificil
    pivot["desempate"] = pivot.meeting_id.map(lambda x: rank(x, "meeting"))
    quarantine = pivot.sort_values(["anio", "score", "desempate"], ascending=[True, False, True]).groupby("anio", sort=True).head(3)
    if len(quarantine) != 33: raise ValueError("Se esperaban tres reuniones por año")
    pool = remaining[remaining.meeting_id.isin(set(quarantine.meeting_id))].copy()
    signals = pool[pool.estrato.ne("")].copy()
    if signals.estrato.value_counts().to_dict() != {"neutral_dificil": 79, "candidato_hawkish": 71, "candidato_dovish": 41}: raise ValueError("Estratos cambiaron")
    selected = signals.copy(); needed = 300 - len(selected)
    random_pool = pool[pool.estrato.eq("")].copy(); random_pool["orden_hash"] = random_pool.intervencion_id.map(lambda x: rank(x, "random"))
    # Relleno equilibrado por reunión, sin usar etiquetas ni predicciones.
    random_pool = random_pool.sort_values(["meeting_id", "orden_hash"]); random_pool["pos_reunion"] = random_pool.groupby("meeting_id").cumcount()
    filler = random_pool.sort_values(["pos_reunion", "orden_hash"]).head(needed).copy(); filler["estrato"] = "aleatorio_contextual"
    selected = pd.concat([selected, filler], ignore_index=True); selected["orden_final"] = selected.intervencion_id.map(lambda x: rank(x, "blind")); selected = selected.sort_values("orden_final").reset_index(drop=True)
    if len(selected) != 300 or selected.intervencion_id.nunique() != 300 or set(selected.intervencion_id) & excluded: raise ValueError("Muestra inválida")
    return selected, quarantine.sort_values(["anio", "meeting_id"]), corpus_path


def workbook(rows, path):
    wb = Workbook(); start = wb.active; start.title = "Inicio"; sheet = wb.create_sheet("Anotación"); guide = wb.create_sheet("Codebook v3")
    navy, yellow, green = "173E4C", "FFF2CC", "DFF1E9"
    start["A1"] = "Evaluación ciega v3 · 300 intervenciones (2005–2015)"; start["A1"].font = Font(size=16, bold=True, color="FFFFFF"); start["A1"].fill = PatternFill("solid", fgColor=navy)
    instructions = ["Este instrumento es evaluación ciega: no contiene predicciones ni etiquetas sugeridas.", "Las intervenciones provienen de 33 reuniones en cuarentena. No intentes completar cuotas H/D/N.", "Lee la unidad completa y aplica dirección monetaria doméstica respaldada del codebook v3.", "Completa etiqueta, relevancia, confianza, cita literal exacta de hasta 300 caracteres y fundamento.", "Usa no_puedo_decidir solo si falta información indispensable y explica qué falta."]
    for i, text in enumerate(instructions, 3): start.cell(i, 1, text).alignment = Alignment(wrap_text=True, vertical="top"); start.row_dimensions[i].height = 34
    start.column_dimensions["A"].width = 120
    headers = ["orden", "intervencion_id", "fecha_reunion", "actor", "cargo", "texto", "etiqueta_v3", "es_relevante_v3", "confianza", "cita_literal", "fundamento", "estado"]
    widths = [8, 31, 15, 28, 26, 100, 19, 19, 14, 56, 56, 18]
    for col, (header, width) in enumerate(zip(headers, widths), 1):
        cell = sheet.cell(1, col, header); cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor=navy); cell.alignment = Alignment(wrap_text=True); sheet.column_dimensions[cell.column_letter].width = width
    for i, row in enumerate(rows.itertuples(index=False), 2):
        values = [i - 1, row.intervencion_id, row.fecha, row.actor, row.cargo, row.texto, "", "", "", "", "", "Pendiente"]
        for col, value in enumerate(values, 1):
            cell = sheet.cell(i, col, value); cell.alignment = Alignment(wrap_text=True, vertical="top")
            if 7 <= col <= 11: cell.fill = PatternFill("solid", fgColor=yellow); cell.protection = Protection(locked=False)
        sheet.cell(i, 12, f'=IF(OR(G{i}="",H{i}="",I{i}=""),"Pendiente",IF(G{i}="no_puedo_decidir",IF(K{i}="","Falta fundamento","Lista"),IF(H{i}="1",IF(OR(J{i}="",LEN(J{i})>300),"Revisar cita",IF(K{i}="","Falta fundamento","Lista")),IF(AND(H{i}="0",G{i}="neutral",K{i}<>""),"Lista","Revisar relevancia"))))')
        sheet.row_dimensions[i].height = 115
    for col, options in (("G", "hawkish,dovish,neutral,no_puedo_decidir"), ("H", "1,0"), ("I", "alta,media,baja")):
        dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True); sheet.add_data_validation(dv); dv.add(f"{col}2:{col}301")
    citation = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="300", allow_blank=True); sheet.add_data_validation(citation); citation.add("J2:J301")
    sheet.freeze_panes = "G2"; sheet.auto_filter.ref = "A1:L301"; sheet.protection.sheet = True; sheet.protection.selectUnlockedCells = True; sheet.protection.selectLockedCells = False
    guide.column_dimensions["A"].width = 120
    for i, line in enumerate((ROOT / "docs/codebook_v3.md").read_text(encoding="utf-8").splitlines(), 1):
        guide.cell(i, 1, line).alignment = Alignment(wrap_text=True, vertical="top")
        if line.startswith("#"): guide.cell(i, 1).font = Font(bold=True, color=navy); guide.cell(i, 1).fill = PatternFill("solid", fgColor=green)
    guide.freeze_panes = "A2"; wb.active = 0; wb.save(path)


def prepare(out=OUT):
    out = Path(out)
    if out.exists(): raise FileExistsError("No sobrescribir evaluación")
    rows, meetings, corpus = select(); out.mkdir(parents=True)
    fields = ["orden", "intervencion_id", "fecha_reunion", "actor", "cargo", "texto", "etiqueta_v3", "es_relevante_v3", "confianza", "cita_literal", "fundamento", "estado"]
    blind = pd.DataFrame([{"orden": i, "intervencion_id": r.intervencion_id, "fecha_reunion": r.fecha, "actor": r.actor, "cargo": r.cargo, "texto": r.texto, "etiqueta_v3": "", "es_relevante_v3": "", "confianza": "", "cita_literal": "", "fundamento": "", "estado": "pendiente"} for i, r in enumerate(rows.itertuples(index=False), 1)], columns=fields)
    blind.to_csv(out / "evaluacion_ciega_300_v3.csv", index=False, lineterminator="\n"); workbook(rows, out / "evaluacion_ciega_300_v3.xlsx")
    meetings[["anio", "meeting_id", "candidato_dovish", "candidato_hawkish", "neutral_dificil"]].to_csv(out / "reuniones_cuarentena.csv", index=False, lineterminator="\n")
    summary = {"filas": 300, "ids_unicos": 300, "periodo": [2005, 2015], "reuniones_cuarentena": 33, "reuniones_por_anio": 3,
        "estratos_recuperacion_ocultos": dict(Counter(rows.estrato)), "advertencia": "estratos de recuperación, no etiquetas",
        "etiquetas_visibles": False, "predicciones_visibles_o_consultadas": False, "rol": "evaluacion ciega challenge agrupada; no train"}
    (out / "resumen.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    protocol = {"seleccion_reuniones": "tres por año por cobertura de señales léxicas prefijadas; sin etiquetas ni predicciones", "cuarentena": "todo dato etiquetado de las 33 reuniones se excluye del ajuste final",
        "seleccion_filas": "todas las 191 señales disponibles más 109 contextuales equilibradas por reunión", "candidato_congelado": "W+C+89+600, parámetros y pesos de caracteres de la ronda preregistrada",
        "apertura": "una vez tras recepción, control y cierre de etiquetas; reportar también por estrato", "limitacion": "las reuniones aparecieron en desarrollo histórico previo, aunque serán excluidas del ajuste prospectivo",
        "fuentes_sha256": {str(path.relative_to(ROOT)): sha(path) for path in [corpus, ROOT / "docs/codebook_v3.md", Path(__file__), *REFERENCE_PATHS, *INSTRUMENT_PATHS]}}
    (out / "protocolo.json").write_text(json.dumps(protocol, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    outputs = ["evaluacion_ciega_300_v3.csv", "evaluacion_ciega_300_v3.xlsx", "reuniones_cuarentena.csv", "resumen.json", "protocolo.json"]
    (out / "manifest.json").write_text(json.dumps({"sha256_salidas": {name: sha(out / name) for name in outputs}}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return summary


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__); parser.add_argument("--salida", type=Path, default=OUT); print(json.dumps(prepare(parser.parse_args().salida), ensure_ascii=False, indent=2))
