"""Prepara 300 intervenciones nuevas enriquecidas para anotación v3, sin etiquetas/modelos."""
import argparse
from collections import Counter
import csv
import hashlib
import json
from pathlib import Path
import re

RAIZ=Path(__file__).resolve().parents[1]
SALIDA=RAIZ/'data/muestras/anotacion_adicional_300_v3'
SEED='anotacion-adicional-v3-20260917'
CUOTAS={'senal_h_sin_d':100,'senal_d_sin_h':100,'senales_mixtas':50,'contexto_monetario':50}
MUESTRAS_PREVIAS=['estrato_enriquecido.csv','estrato_fases.csv','estrato_tanda9.csv',
                  'gold_ciego_300.csv','piloto_300.csv','test_retest_30.csv']
PATRON_H=re.compile(r'\b(subir|subida|alza|alzas|elevar|aumentar|incrementar|endurec|contractiv|restrictiv|retir(?:ar|o|ando)|normaliz|sesgo al alza)\w*',re.I)
PATRON_D=re.compile(r'\b(bajar|baja|bajas|reducir|reducci|recort|relaj|expansiv|flexibili|estimulo|estímulo|sesgo a la baja)\w*',re.I)
PATRON_M=re.compile(r'\b(tpm|tasa de pol[ií]tica|pol[ií]tica monetaria|tasa de instancia|facilidad de liquidez|estimulo monetario|estímulo monetario)\b',re.I)


def sha256(ruta):return hashlib.sha256(Path(ruta).read_bytes()).hexdigest()
def clave_texto(x):return ' '.join(x.split()).lower()
def orden_hash(x):return hashlib.sha256((SEED+'|'+x).encode()).hexdigest()

def guardar_json(ruta,obj):
    with Path(ruta).open('x',encoding='utf-8') as f:json.dump(obj,f,ensure_ascii=False,indent=2);f.write('\n')


def estrato(texto):
    h,d=bool(PATRON_H.search(texto)),bool(PATRON_D.search(texto))
    if h and not d:return 'senal_h_sin_d'
    if d and not h:return 'senal_d_sin_h'
    if h and d:return 'senales_mixtas'
    return 'contexto_monetario'


def seleccionar():
    corpus_path=RAIZ/'data/L0/corpus.csv';ref_path=RAIZ/'data/evaluacion/referencia_v3/referencia_v3.csv'
    with corpus_path.open(encoding='utf-8',newline='') as f:corpus=list(csv.DictReader(f))
    with ref_path.open(encoding='utf-8',newline='') as f:ref=list(csv.DictReader(f))
    revisados={x['intervencion_id'] for x in ref};por_id={x['intervencion_id']:x for x in corpus}
    previamente_muestreados=set()
    for nombre in MUESTRAS_PREVIAS:
        ruta=RAIZ/'data/muestras'/nombre
        with ruta.open(encoding='utf-8',newline='') as f:
            lector=csv.DictReader(f)
            previamente_muestreados.update(x['intervencion_id'] for x in lector)
    excluir_ids=revisados|previamente_muestreados
    textos_revisados={clave_texto(por_id[x]['texto']) for x in excluir_ids if x in por_id}
    candidatos=[]
    for x in corpus:
        if x['intervencion_id'] in excluir_ids or clave_texto(x['texto']) in textos_revisados:continue
        if x['flag_texto_danado']=='True' or not PATRON_M.search(x['texto']):continue
        y=dict(x);y['_estrato']=estrato(x['texto']);candidatos.append(y)
    por_estrato={e:sorted([x for x in candidatos if x['_estrato']==e],key=lambda x:orden_hash(x['intervencion_id'])) for e in CUOTAS}
    cursores=Counter();elegidos=[];conteo_reunion=Counter();conteo_anio=Counter();pendientes=dict(CUOTAS)
    # Ciclo equilibrado; topes impiden que una reunión o año domine el instrumento.
    while sum(pendientes.values()):
        progreso=False
        for e in CUOTAS:
            if not pendientes[e]:continue
            lista=por_estrato[e]
            while cursores[e]<len(lista):
                x=lista[cursores[e]];cursores[e]+=1
                if conteo_reunion[x['meeting_id']]>=4 or conteo_anio[x['anio']]>=32:continue
                if clave_texto(x['texto']) in {clave_texto(z['texto']) for z in elegidos}:continue
                elegidos.append(x);conteo_reunion[x['meeting_id']]+=1;conteo_anio[x['anio']]+=1
                pendientes[e]-=1;progreso=True;break
        if not progreso:raise ValueError('No se completaron cuotas con los topes fijados')
    elegidos.sort(key=lambda x:orden_hash('orden|'+x['intervencion_id']))
    if len(elegidos)!=300 or len({x['intervencion_id'] for x in elegidos})!=300:raise ValueError('Muestra incompleta')
    if set(x['intervencion_id'] for x in elegidos)&excluir_ids:raise ValueError('ID revisado o muestreado previamente')
    return elegidos,corpus_path,ref_path,conteo_reunion,conteo_anio


def crear_excel(filas,ruta,codebook):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment,Font,PatternFill,Protection
    from openpyxl.worksheet.datavalidation import DataValidation
    wb=Workbook();inicio=wb.active;inicio.title='Inicio';hoja=wb.create_sheet('Anotación');guia=wb.create_sheet('Codebook v3')
    azul='173E4C';amarillo='FFF2CC';verde='DFF1E9'
    inicio['A1']='Anotación adicional v3 · 300 intervenciones';inicio['A1'].font=Font(size=16,bold=True,color='FFFFFF');inicio['A1'].fill=PatternFill('solid',fgColor=azul)
    instrucciones=[
      'Estas 300 intervenciones no tienen etiquetas IA ni predicciones visibles. No intentes completar cuotas H/D/N.',
      'Lee la intervención completa y aplica dirección monetaria doméstica respaldada del codebook v3.',
      'Completa etiqueta, relevancia, confianza, cita literal de hasta 300 caracteres y fundamento.',
      'Si no puedes decidir, usa no_puedo_decidir y explica qué información falta. No lo conviertas en neutral por defecto.',
      'Esta muestra amplía datos reales, pero comparte reuniones históricas con el train y no será el test final independiente.' ]
    for i,t in enumerate(instrucciones,3):inicio.cell(i,1,t);inicio.cell(i,1).alignment=Alignment(wrap_text=True,vertical='top');inicio.row_dimensions[i].height=35
    inicio.column_dimensions['A'].width=115
    headers=['orden','intervencion_id','fecha_reunion','actor','cargo','texto','etiqueta_v3','es_relevante_v3','confianza','cita_literal','fundamento','estado']
    for j,h in enumerate(headers,1):
        c=hoja.cell(1,j,h);c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor=azul);c.alignment=Alignment(wrap_text=True)
    widths=[8,30,14,28,25,100,18,18,14,55,55,18]
    for j,w in enumerate(widths,1):hoja.column_dimensions[chr(64+j)].width=w
    for i,x in enumerate(filas,2):
        vals=[i-1,x['intervencion_id'],x['fecha'],x['actor'],x['cargo'],x['texto'],'','','','','','Pendiente']
        for j,v in enumerate(vals,1):
            c=hoja.cell(i,j,v);c.alignment=Alignment(wrap_text=True,vertical='top')
            if 7<=j<=11:c.fill=PatternFill('solid',fgColor=amarillo);c.protection=Protection(locked=False)
        hoja.cell(i,12,f'=IF(OR(G{i}="",H{i}="",I{i}=""),"Pendiente",IF(G{i}="no_puedo_decidir",IF(K{i}="","Falta fundamento","Lista"),IF(H{i}="1",IF(OR(J{i}="",LEN(J{i})>300),"Revisar cita",IF(K{i}="","Falta fundamento","Lista")),IF(AND(H{i}="0",G{i}="neutral",K{i}<>""),"Lista","Revisar relevancia"))))')
        hoja.row_dimensions[i].height=115
    for col,opciones in [('G','hawkish,dovish,neutral,no_puedo_decidir'),('H','1,0'),('I','alta,media,baja')]:
        dv=DataValidation(type='list',formula1='"'+opciones+'"',allow_blank=True);hoja.add_data_validation(dv);dv.add(f'{col}2:{col}301')
    cita=DataValidation(type='textLength',operator='lessThanOrEqual',formula1='300',allow_blank=True);hoja.add_data_validation(cita);cita.add('J2:J301')
    hoja.freeze_panes='G2';hoja.auto_filter.ref='A1:L301';hoja.protection.sheet=True;hoja.protection.selectUnlockedCells=True;hoja.protection.selectLockedCells=False
    guia.column_dimensions['A'].width=120
    for i,linea in enumerate(codebook.splitlines(),1):
        guia.cell(i,1,linea);guia.cell(i,1).alignment=Alignment(wrap_text=True,vertical='top')
        if linea.startswith('#'):guia.cell(i,1).font=Font(bold=True,color=azul);guia.cell(i,1).fill=PatternFill('solid',fgColor=verde)
    guia.freeze_panes='A2';wb.active=0;wb.save(ruta)


def preparar(salida=SALIDA):
    salida=Path(salida);nombres=['anotacion_adicional_300_v3.csv','anotacion_adicional_300_v3.xlsx','resumen.json','protocolo.json','manifest.json']
    if salida.exists():raise FileExistsError('No sobrescribir muestra adicional')
    filas,corpus,ref,reuniones,anios=seleccionar();salida.mkdir(parents=True)
    campos=['orden','intervencion_id','fecha_reunion','actor','cargo','texto','etiqueta_v3','es_relevante_v3','confianza','cita_literal','fundamento','estado']
    salida_csv=[]
    for i,x in enumerate(filas,1):salida_csv.append({'orden':i,'intervencion_id':x['intervencion_id'],'fecha_reunion':x['fecha'],'actor':x['actor'],'cargo':x['cargo'],'texto':x['texto'],'etiqueta_v3':'','es_relevante_v3':'','confianza':'','cita_literal':'','fundamento':'','estado':'pendiente'})
    with (salida/nombres[0]).open('x',encoding='utf-8',newline='') as f:w=csv.DictWriter(f,fieldnames=campos,lineterminator='\n');w.writeheader();w.writerows(salida_csv)
    codebook=(RAIZ/'docs/codebook_v3.md').read_text(encoding='utf-8');crear_excel(filas,salida/nombres[1],codebook)
    resumen={'filas':300,'ids_unicos':300,'cuotas_seleccion_sin_etiquetas':CUOTAS,'reuniones':len(reuniones),'max_por_reunion':max(reuniones.values()),'anios':dict(sorted(anios.items())),'max_por_anio':max(anios.values()),'etiquetas_visibles':False,'predicciones_visibles':False,'respuestas_pendientes':300,'rol':'ampliacion_real_anotada; no test final independiente'}
    guardar_json(salida/nombres[2],resumen)
    fuentes=[corpus,ref,RAIZ/'docs/codebook_v3.md',Path(__file__)]+[RAIZ/'data/muestras'/n for n in MUESTRAS_PREVIAS]
    guardar_json(salida/nombres[3],{'operacion':'Muestreo determinista enriquecido por vocabulario, sin consultar etiquetas o predicciones.','cuotas':CUOTAS,'topes':{'por_reunion':4,'por_anio':32},'muestras_previas_excluidas':MUESTRAS_PREVIAS,'limitacion':'Todas las reuniones de L0 ya aparecen en IA-base; esta muestra no puede ser test independiente por reunión.','fuentes_sha256':{str(p.relative_to(RAIZ)):sha256(p) for p in fuentes}})
    guardar_json(salida/nombres[4],{'sha256_salidas':{n:sha256(salida/n) for n in nombres if n!='manifest.json'}})
    return resumen


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--salida',type=Path,default=SALIDA)
    print(json.dumps(preparar(p.parse_args().salida),ensure_ascii=False,indent=2))
