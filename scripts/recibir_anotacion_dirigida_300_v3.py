"""Recibe y audita estructuralmente la segunda entrega dirigida de 300 casos v3."""
import csv
import hashlib
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data/recepcion/anotacion_dirigida_300_v3_v1"
SOURCE = OUT / "anotacion_original.xlsx"
BLIND = ROOT / "data/muestras/anotacion_dirigida_300_v3_v1/anotacion_dirigida_300_v3.csv"
LABELS = {"H": "hawkish", "D": "dovish", "N": "neutral", "hawkish": "hawkish", "dovish": "dovish", "neutral": "neutral"}
MANUAL = {"RPM-2008-05-08:1817:1": ("la evaluación del Consejo, apoyado en otras medidas", 172)}


def norm(value): return " ".join(str("" if value is None else value).split())
def sha(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def repair_quote(identity, quote, text):
    q, t = norm(quote), norm(text)
    if q in t: return q, "sin_cambio"
    start = t.casefold().find(q.casefold())
    if start >= 0: return t[start:start + len(q)], "capitalizacion_repuesta_desde_texto"
    trimmed = q[:-3].rstrip() if q.endswith("...") else q
    start = t.casefold().find(trimmed.casefold())
    if start >= 0: return t[start:start + len(trimmed)], "elipsis_no_literal_retirada"
    if identity in MANUAL:
        marker, length = MANUAL[identity]; start = t.casefold().find(marker.casefold())
        if start < 0: raise ValueError(f"Marcador ausente: {identity}")
        return t[start:start + length].rstrip(), "transcripcion_repuesta_desde_texto"
    raise ValueError(f"Cita no reparable exactamente: {identity}")


def read(path):
    with Path(path).open(encoding="utf-8", newline="") as f: return list(csv.DictReader(f))


def write_csv(path, rows, fields=None):
    fields = fields or list(rows[0])
    with path.open("x", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, lineterminator="\n"); writer.writeheader(); writer.writerows(rows)


def run(out=OUT):
    out = Path(out); names = ["anotaciones_canonicas.csv", "incidencias_reparadas.csv", "revision_semantica_300.csv", "resumen.json", "protocolo.json", "manifest.json"]
    if any((out / name).exists() for name in names): raise FileExistsError("No sobrescribir recepción")
    blind = {r["intervencion_id"]: r for r in read(BLIND)}
    wb = load_workbook(SOURCE, read_only=True, data_only=False)
    if wb.sheetnames != ["Inicio", "Anotación", "Codebook v3"]: raise ValueError("Hojas inesperadas")
    ws = wb["Anotación"]; headers = [cell.value for cell in ws[1]]
    expected = ["orden", "intervencion_id", "fecha_reunion", "actor", "cargo", "texto", "etiqueta_v3", "es_relevante_v3", "confianza", "cita_literal", "fundamento", "estado"]
    if headers != expected or ws.max_row != 301: raise ValueError("Esquema inesperado")
    human = [dict(zip(headers, values)) for values in ws.iter_rows(min_row=2, max_row=301, values_only=True)]
    if len({r["intervencion_id"] for r in human}) != 300 or set(blind) != {r["intervencion_id"] for r in human}: raise ValueError("IDs distintos del instrumento")
    canonical, incidents = [], []
    for row in human:
        identity = norm(row["intervencion_id"]); original = blind[identity]
        for field in ["orden", "fecha_reunion", "actor", "cargo", "texto"]:
            if norm(row[field]) != norm(original[field]): raise ValueError(f"Campo fuente modificado: {identity}/{field}")
        raw_label = norm(row["etiqueta_v3"]); relevance = norm(row["es_relevante_v3"]); confidence = norm(row["confianza"]); rationale = norm(row["fundamento"])
        if raw_label not in LABELS or relevance not in {"0", "1"} or confidence not in {"alta", "media", "baja"} or not rationale: raise ValueError(f"Respuesta inválida: {identity}")
        label = LABELS[raw_label]
        if relevance == "0" and label != "neutral": raise ValueError(f"Fila irrelevante direccional: {identity}")
        raw_quote = norm(row["cita_literal"])
        if relevance == "1":
            if not raw_quote: raise ValueError(f"Falta cita: {identity}")
            quote, reason = repair_quote(identity, raw_quote, row["texto"])
            if len(quote) > 300 or quote not in norm(row["texto"]): raise ValueError(f"Cita final inválida: {identity}")
            if reason != "sin_cambio": incidents.append({"intervencion_id": identity, "tipo": reason, "cita_original": raw_quote, "cita_reparada": quote})
        else: quote, reason = "", "sin_cambio"
        canonical.append({"orden": int(row["orden"]), "intervencion_id": identity, "fecha_reunion": norm(row["fecha_reunion"]), "actor": norm(row["actor"]), "cargo": norm(row["cargo"]), "texto": norm(row["texto"]),
            "etiqueta_humana_original": raw_label, "etiqueta_v3": label, "es_relevante_v3": relevance, "confianza_humana": confidence,
            "cita_literal": quote, "fundamento_humano": rationale, "estado_original": norm(row["estado"]), "estado_estructural": "valido",
            "estado_semantico": "pendiente_segunda_revision", "sha256_texto": hashlib.sha256(norm(row["texto"]).encode()).hexdigest()})
    out.mkdir(parents=True, exist_ok=True)
    write_csv(out / names[0], canonical)
    write_csv(out / names[1], incidents, ["intervencion_id", "tipo", "cita_original", "cita_reparada"])
    review_fields = ["intervencion_id", "fecha_reunion", "actor", "cargo", "texto", "etiqueta_v3", "es_relevante_v3", "confianza_humana", "cita_literal", "fundamento_humano", "revision_etiqueta_v3", "revision_relevancia_v3", "revision_confianza", "revision_cita_literal", "revision_fundamento", "revision_estado"]
    review = []
    for row in canonical:
        item = {key: row.get(key, "") for key in review_fields}; item.update({key: "" for key in review_fields if key.startswith("revision_")}); review.append(item)
    write_csv(out / names[2], review, review_fields)
    summary = {"filas": 300, "ids_unicos": 300, "etiquetas_v3": dict(Counter(r["etiqueta_v3"] for r in canonical)), "relevancia_v3": dict(Counter(r["es_relevante_v3"] for r in canonical)),
        "confianza": dict(Counter(r["confianza_humana"] for r in canonical)), "citas_reparadas": len(incidents), "campos_fuente_modificados": 0,
        "filas_estructuralmente_validas": 300, "filas_semanticamente_cerradas": 0, "cola_segunda_revision": 300, "lista_para_entrenamiento": False}
    (out / names[3]).write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    protocol = {"origen_github": "main@ca9b80f:anotacion_dirigida_300_v3.xlsx", "sha256_origen": sha(SOURCE), "sha256_muestra_ciega": sha(BLIND),
        "normalizacion_etiquetas": {"H": "hawkish", "D": "dovish", "N": "neutral"}, "reparacion_citas": "solo reextracción literal del texto fuente",
        "regla_revision": "segunda revisión de las 300 filas bajo dirección propia respaldada, sin predicciones", "entrenamiento_permitido": False}
    (out / names[4]).write_text(json.dumps(protocol, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (out / names[5]).write_text(json.dumps({"sha256_salidas": {name: sha(out / name) for name in names[:-1]}}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return summary


if __name__ == "__main__": print(json.dumps(run(), ensure_ascii=False, indent=2))
