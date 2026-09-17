"""Excel editable de los mismos 30 casos y descarga HTTP real (sin Blob).

No lee claves IA ni respuestas humanas: utiliza solo el payload público del
formulario. Conserva la interfaz anterior antes de añadir el enlace XLSX.
"""
# ---- 1. Configuración, estilo y utilidades de escritura segura ----
import argparse
from datetime import datetime, timezone
from functools import partial
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import json
import math
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.workbook.properties import CalcProperties

from utilidades import cargar_script, exigir_salidas_nuevas, sha256

ANTERIOR = cargar_script('28_reparar_exportacion_revision.py')
SALIDA = ANTERIOR.SALIDA
NOMBRE_EXCEL = 'revision_entrenamiento_30.xlsx'
VERSION = 'revision_excel_v1'
AZUL = '173E4C'
VERDE = 'DFF1E9'
AMARILLO = 'FFF2CC'
GRIS = 'EDF1F3'


def texto(celda,valor):
    """Un texto que empieza por = no debe convertirse en una fórmula ejecutable."""
    celda.value = str(valor)
    celda.data_type = 's'
    celda.font = Font(name='Calibri',size=11,color='173E4C')
    celda.alignment = Alignment(wrap_text=True,vertical='top')
    return celda


def titulo(hoja,fila,valor,ultima='H'):
    hoja.merge_cells(f'A{fila}:{ultima}{fila}')
    celda = texto(hoja[f'A{fila}'],valor)
    celda.fill = PatternFill('solid',fgColor=AZUL)
    celda.font = Font(name='Calibri',size=17,bold=True,color='FFFFFF')
    hoja.row_dimensions[fila].height = 34


def enlace(celda,etiqueta,hoja,destino):
    texto(celda,etiqueta)
    celda.hyperlink = Hyperlink(ref=celda.coordinate,location=f"'{hoja}'!{destino}",display=etiqueta)
    celda.font = Font(name='Calibri',size=11,color='147D70',underline='single',bold=True)


def editable(celda):
    celda.fill = PatternFill('solid',fgColor=AMARILLO)
    celda.alignment = Alignment(wrap_text=True,vertical='top')
    celda.protection = Protection(locked=False)


def lista(hoja,rango,opciones):
    regla = DataValidation(type='list',formula1='"'+','.join(opciones)+'"',allow_blank=True)
    regla.showDropDown = False
    regla.showErrorMessage = True
    regla.errorStyle = 'stop'
    regla.errorTitle = 'Usa una opción de la lista'
    regla.error = 'Selecciona una opción del desplegable o deja la celda pendiente.'
    hoja.add_data_validation(regla)
    regla.add(rango)


def proteger(hoja):
    # Protección contra cambios accidentales, sin contraseña; no es seguridad fuerte.
    hoja.protection.sheet = True
    hoja.protection.selectLockedCells = False
    hoja.protection.selectUnlockedCells = False
    hoja.protection.formatColumns = False
    hoja.protection.formatRows = False
    hoja.protection.autoFilter = False
    hoja.sheet_view.showGridLines = False


# ---- 2. Textos extensos: bloques legibles sin perder un solo carácter ----
def bloques(texto_original,limite=650):
    assert limite>0
    partes=[]
    inicio=0
    while inicio<len(texto_original):
        fin=min(inicio+limite,len(texto_original))
        saltos=[i for i in range(inicio,fin) if texto_original[i]=='\n']
        if len(saltos)>=8:
            fin=saltos[7]+1
        elif fin<len(texto_original):
            corte=texto_original.rfind(' ',inicio,fin)
            if corte>inicio+limite//2:
                fin=corte+1
        partes.append(texto_original[inicio:fin])
        inicio=fin
    assert ''.join(partes)==texto_original
    return partes


def formula_estado(fila):
    c,d,e,f=(f'{col}{fila}' for col in 'CDEF')
    return (f'=IF(LEN(TRIM({c}))=0,"Pendiente",IF({c}="No puedo decidir",'
        f'IF(LEN(TRIM({f}))=0,"Falta motivo","Lista"),'
        f'IF(OR({c}="Hawkish",{c}="Dovish",{c}="Neutral"),'
        f'IF({d}="No",IF({c}<>"Neutral","Revisar relevancia",IF(LEN(TRIM({f}))=0,"Falta motivo","Lista")),'
        f'IF({d}<>"Sí","Falta relevancia",IF(LEN(TRIM({e}))=0,"Falta cita",IF(LEN({e})>300,"Cita >300","Lista")))),"Revisar postura")))')


def crear_libro(publico,ruta):
    exigir_salidas_nuevas(ruta)
    assert len(publico['casos'])==30
    libro=Workbook()
    inicio=libro.active
    inicio.title='Inicio'
    respuestas=libro.create_sheet('Respuestas')
    textos=libro.create_sheet('Textos')
    guia=libro.create_sheet('Guía')
    libro.properties.creator='Proyecto D&H'
    libro.properties.title='Revisión humana de 30 intervenciones de entrenamiento'
    libro.properties.description='Misma muestra; sin etiquetas IA ni respuestas recuperadas del navegador.'
    libro.calculation=CalcProperties(calcMode='auto',fullCalcOnLoad=True)
    inicio.column_dimensions['A'].width=32
    inicio.column_dimensions['B'].width=105
    titulo(inicio,1,'Revisión de entrenamiento · 30 casos','B')
    instrucciones=[
        '1. Guarda una copia en tu equipo antes de empezar y usa Ctrl+S / Cmd+S durante la revisión.',
        '2. En Respuestas, completa solo las celdas amarillas. Las decisiones están vacías: este archivo no recupera el avance perdido del navegador.',
        '3. Usa «Leer texto» para ir a la intervención completa. En Textos, pulsa el código R01–R30 para volver a su respuesta.',
        '4. Decide H/D/N según la guía de la TPM chilena. «No puedo decidir» registra una duda, no una cuarta clase para entrenar.',
        '5. Relevante: Sí salvo formalidad/logística o texto no evaluable. Si marcas No, usa Neutral y explica el motivo.',
        '6. Para H/D/N relevantes, pega una cita literal de hasta 300 caracteres. Para una duda o contenido no evaluable, explica el motivo.',
        '7. Estado revisa campos requeridos, no la corrección de tu criterio ni si la cita es literal. No indica acuerdo con la IA.',
        '8. Guarda el XLSX y adjúntalo al chat, aunque sea parcial. No hace falta volver a la web ni convertir a JSON.']
    for fila,instruccion in enumerate(instrucciones,3):
        inicio.merge_cells(start_row=fila,start_column=1,end_row=fila,end_column=2)
        texto(inicio.cell(fila,1),instruccion)
        inicio.row_dimensions[fila].height=44
    enlace(inicio['A12'],'Ir a mis respuestas','Respuestas','A7')
    enlace(inicio['B12'],'Leer la guía vigente','Guía','A1')
    for fila,rotulo in [(14,'Nombre (opcional)'),(15,'Fecha de revisión (opcional)'),(16,'Ayuda recibida'),(17,'¿Viste etiquetas IA anteriores?')]:
        texto(inicio.cell(fila,1),rotulo)
        editable(inicio.cell(fila,2))
        inicio.row_dimensions[fila].height=32
    lista(inicio,'B16',['Ninguna','IA','Otra persona','IA y persona','Prefiero no declarar'])
    lista(inicio,'B17',['No','Sí','No recuerdo'])
    texto(inicio['A19'],'Respuestas con campos completos')
    inicio['B19']='=COUNTIF(Respuestas!G7:G36,"Lista")'
    texto(inicio['A21'],'Versión de la muestra')
    texto(inicio['B21'],publico['version'])
    texto(inicio['A22'],'Huella de la muestra (no editar)')
    texto(inicio['B22'],publico['muestra_sha256'])
    inicio.row_dimensions[22].height=36
    texto(inicio['A24'],'Independencia y confidencialidad')
    texto(inicio['B24'],'Responde antes de consultar otra IA o etiquetas anteriores. No intentes completar cuotas de clases. No hay etiquetas IA ocultas en el libro; declara si recibiste ayuda.')
    inicio.row_dimensions[24].height=50
    inicio.freeze_panes='B3'

    titulo(respuestas,1,'Tus decisiones · completar las celdas amarillas')
    respuestas.merge_cells('A2:H2')
    texto(respuestas['A2'],'Las etiquetas IA no se muestran. Puedes guardar y devolver un avance parcial; no inventes una decisión para completar la hoja.')
    respuestas.row_dimensions[2].height=34
    respuestas.merge_cells('A3:H3')
    texto(respuestas['A3'],'C–F son editables. Usa el enlace de la última columna para leer cada texto completo. «Lista» solo significa que completaste los campos requeridos.')
    respuestas.row_dimensions[3].height=32
    enlace(respuestas['A4'],'Instrucciones','Inicio','A1')
    encabezados=['Caso','Actor','Mi postura','Relevante','Cita literal (máx. 300)','Motivo / duda','Estado','Texto completo']
    for columna,(rotulo,ancho) in enumerate(zip(encabezados,[10,28,23,13,52,42,22,19]),1):
        celda=texto(respuestas.cell(6,columna),rotulo)
        celda.fill=PatternFill('solid',fgColor=AZUL)
        celda.font=Font(name='Calibri',size=11,bold=True,color='FFFFFF')
        respuestas.column_dimensions[celda.column_letter].width=ancho
    respuestas.row_dimensions[6].height=32
    respuestas.freeze_panes='C7'
    respuestas.auto_filter.ref='A6:H36'
    lista(respuestas,'C7:C36',['Hawkish','Dovish','Neutral','No puedo decidir'])
    lista(respuestas,'D7:D36',['Sí','No'])
    cita=DataValidation(type='textLength',operator='lessThanOrEqual',formula1=300,allow_blank=True)
    cita.showErrorMessage=True
    cita.errorStyle='stop'
    cita.errorTitle='Cita demasiado larga'
    cita.error='Usa un fragmento literal de hasta 300 caracteres.'
    respuestas.add_data_validation(cita)
    cita.add('E7:E36')
    respuestas.conditional_formatting.add('G7:G36',CellIsRule(operator='equal',formula=['"Lista"'],fill=PatternFill('solid',fgColor=VERDE)))
    respuestas.conditional_formatting.add('G7:G36',FormulaRule(formula=['AND(G7<>"Lista",G7<>"Pendiente")'],fill=PatternFill('solid',fgColor='FCE4D6')))

    titulo(textos,1,'Intervenciones completas · sin etiquetas IA','C')
    textos.merge_cells('A2:C2')
    texto(textos['A2'],'Lee todos los bloques del caso. El texto solo se divide para caber en la pantalla: no está resumido ni recortado. Pulsa su código para volver a Respuestas.')
    textos.row_dimensions[2].height=42
    for col,ancho in [('A',11),('B',13),('C',112)]: textos.column_dimensions[col].width=ancho
    textos.freeze_panes='C4'
    ubicaciones={}
    fila_texto=4
    for fila,caso in enumerate(publico['casos'],7):
        texto(respuestas.cell(fila,1),caso['caso_id'])
        texto(respuestas.cell(fila,2),caso['actor'])
        for columna in range(3,7): editable(respuestas.cell(fila,columna))
        respuestas.cell(fila,7,formula_estado(fila))
        respuestas.cell(fila,7).alignment=Alignment(wrap_text=True,vertical='top')
        respuestas.cell(fila,7).fill=PatternFill('solid',fgColor=GRIS)
        respuestas.row_dimensions[fila].height=98
        enlace(respuestas.cell(fila,8),'Leer '+caso['caso_id'],'Textos',f'A{fila_texto}')
        enlace(textos.cell(fila_texto,1),caso['caso_id']+' ↩','Respuestas',f'A{fila}')
        texto(textos.cell(fila_texto,3),' · '.join([caso['actor'],caso['cargo'],caso['fecha']]))
        for columna in range(1,4): textos.cell(fila_texto,columna).fill=PatternFill('solid',fgColor=VERDE)
        textos.row_dimensions[fila_texto].height=50
        ubicaciones[caso['caso_id']]={'fila_respuesta':fila,'filas_texto':[]}
        partes=bloques(caso['texto'])
        for i,parte in enumerate(partes,1):
            fila_texto+=1
            texto(textos.cell(fila_texto,1),caso['caso_id'])
            texto(textos.cell(fila_texto,2),f'{i}/{len(partes)}')
            texto(textos.cell(fila_texto,3),parte)
            lineas=sum(max(1,math.ceil(len(linea)/90)) for linea in parte.split('\n'))
            altura=max(28,lineas*17+10)
            assert altura<=409, 'Bloque demasiado alto para Excel'
            textos.row_dimensions[fila_texto].height=altura
            ubicaciones[caso['caso_id']]['filas_texto'].append(fila_texto)
        fila_texto+=2
    titulo(guia,1,'Guía v2 · secciones 1–3','B')
    guia.column_dimensions['A'].width=4
    guia.column_dimensions['B'].width=125
    enlace(guia['B2'],'Volver a respuestas','Respuestas','A7')
    for fila,linea in enumerate(publico['guia'].split('\n'),4):
        celda=texto(guia.cell(fila,2),linea)
        if linea.startswith('#'):
            celda.font=Font(name='Calibri',size=12,bold=True,color=AZUL)
            celda.fill=PatternFill('solid',fgColor=VERDE)
        guia.row_dimensions[fila].height=max(20,math.ceil(max(1,len(linea))/100)*17+8)
    guia.freeze_panes='B4'
    for hoja in libro:
        proteger(hoja)
        hoja.sheet_properties.pageSetUpPr.fitToPage=True
        hoja.page_setup.orientation='landscape'
        hoja.page_setup.paperSize=hoja.PAPERSIZE_A4
        hoja.page_setup.fitToWidth=1
        hoja.page_setup.fitToHeight=0
    respuestas.print_title_rows='1:6'
    textos.print_title_rows='1:2'
    libro.active=0
    libro.save(ruta)
    return ubicaciones


# ---- 3. Enlace web versionado, sin modificar JavaScript, muestra o autoguardado ----
def actualizar_html(html):
    assert '<!-- revision-excel-v1 -->' not in html
    bloque=f'''<!-- revision-excel-v1 -->
<section class="intro" aria-labelledby="titulo-excel"><h2 id="titulo-excel">Responder en Excel</h2>
<p>Descarga los mismos 30 casos, guarda el archivo en tu equipo y complétalo a tu ritmo. Después adjunta el XLSX en el chat; no necesitas exportar JSON.</p>
<a class="primary" style="display:inline-block;padding:10px 16px;border-radius:9px;text-decoration:none" href="{NOMBRE_EXCEL}" download="{NOMBRE_EXCEL}" target="_blank" rel="noopener">Descargar Excel para completar (.xlsx)</a>
<p class="muted">Incluye textos completos y columnas vacías para tus decisiones, sin etiquetas IA. No recupera respuestas perdidas ni incorpora lo escrito en el navegador. Si el visor bloquea la descarga, usa el archivo adjunto directamente en el chat.</p></section>'''
    assert html.count('<main>')==1
    nuevo=html.replace('<main>','<main>'+bloque,1)
    assert ANTERIOR.payload(nuevo)==ANTERIOR.payload(html)
    return nuevo


def preparar(salida=SALIDA):
    salida=Path(salida)
    pagina=salida/'formulario/index.html'
    manifest=salida/'manifest.json'
    archivo=salida/'archivo_interfaz_v1_1'
    registro=salida/'excel_revision_v1.json'
    excel=salida/NOMBRE_EXCEL
    exigir_salidas_nuevas(excel,archivo,registro)
    anterior=json.loads(manifest.read_text())
    assert sha256(pagina)==anterior['sha256_salidas']['formulario/index.html']
    html=pagina.read_text()
    publico=ANTERIOR.payload(html)
    nuevo=actualizar_html(html)
    ubicaciones=crear_libro(publico,excel)
    archivo.mkdir()
    (archivo/'index.html').write_bytes(pagina.read_bytes())
    (archivo/'manifest.json').write_bytes(manifest.read_bytes())
    ANTERIOR.escribir_atomico(pagina,nuevo)
    anterior['sha256_salidas']['formulario/index.html']=sha256(pagina)
    anterior['sha256_salidas'][NOMBRE_EXCEL]=sha256(excel)
    anterior['interfaz']='v1.2: descarga XLSX; versiones previas conservadas fuera de rutas públicas.'
    ANTERIOR.escribir_atomico(manifest,json.dumps(anterior,ensure_ascii=False,indent=2)+'\n')
    registro.write_text(json.dumps({'version':VERSION,'fecha_utc':datetime.now(timezone.utc).isoformat(),
        'muestra_sha256':publico['muestra_sha256'],'sha256_excel':sha256(excel),
        'sha256_html_previo':sha256(archivo/'index.html'),'sha256_html_actual':sha256(pagina),
        'sha256_generador':sha256(Path(__file__)),'ubicaciones':ubicaciones,
        'sha256_tests':sha256(RAIZ_TEST),'n_casos':30,'n_hojas':4,
        'etiquetas_ia_o_respuestas_navegador_incluidas':False,
        'estado':'Libro para completar; sin decisiones nuevas recibidas ni correcciones.'},ensure_ascii=False,indent=2)+'\n')
    print('Excel creado con 30 casos, textos completos y decisiones vacías:',excel)


# ---- 4. Descarga por HTTP con lista explícita: nunca servir la raíz del proyecto ----
class DescargaRevision(BaseHTTPRequestHandler):
    def __init__(self,*args,directorio=SALIDA,**kwargs):
        self.directorio=Path(directorio)
        super().__init__(*args,**kwargs)

    def responder(self,cuerpo):
        ruta=urlsplit(self.path).path
        if ruta in ('/','/index.html'):
            archivo=self.directorio/'formulario/index.html'
            tipo='text/html; charset=utf-8'
        elif ruta=='/'+NOMBRE_EXCEL:
            archivo=self.directorio/NOMBRE_EXCEL
            tipo='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        else:
            self.send_error(404)
            return
        if not archivo.is_file():
            self.send_error(404)
            return
        self.send_response(200)
        self.send_header('Content-Type',tipo)
        self.send_header('Content-Length',str(archivo.stat().st_size))
        self.send_header('Cache-Control','no-store')
        self.send_header('X-Content-Type-Options','nosniff')
        if ruta=='/'+NOMBRE_EXCEL:
            self.send_header('Content-Disposition',f'attachment; filename="{NOMBRE_EXCEL}"')
        self.end_headers()
        if cuerpo:
            self.wfile.write(archivo.read_bytes())

    def do_GET(self):
        self.responder(True)

    def do_HEAD(self):
        self.responder(False)


RAIZ_TEST=ANTERIOR.RAIZ/'tests/test_excel_revision.py'


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    accion=parser.add_mutually_exclusive_group(required=True)
    accion.add_argument('--preparar',action='store_true')
    accion.add_argument('--servir',action='store_true')
    parser.add_argument('--salida',type=Path,default=SALIDA)
    parser.add_argument('--puerto',type=int,default=8080)
    args=parser.parse_args()
    if args.preparar:
        preparar(args.salida)
    else:
        servidor=ThreadingHTTPServer(('0.0.0.0',args.puerto),partial(DescargaRevision,directorio=args.salida))
        print(f'Revisión y descarga Excel en puerto {args.puerto}; solo dos recursos públicos.',flush=True)
        try:
            servidor.serve_forever()
        finally:
            servidor.server_close()


if __name__=='__main__':
    main()
