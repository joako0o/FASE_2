# preparar_gold_xlsx.py
# =============================================================================
# Libro Excel del gold ciego para etiquetar SIN pelear con codificación ni
# separadores (el .csv canónico es UTF-8 con comas; Excel en español lo abre
# como ANSI con ';' y se ve "sucio": mojibake + todo en una columna).
#
# QUÉ HACE
# Lee el canónico data/muestras/gold_ciego_300.csv (NO lo modifica) y escribe
# data/muestras/gold_ciego_300.xlsx con:
#   - Hoja "etiquetar": las 306 filas, 16 columnas, cabecera congelada,
#     autofiltro, anchos ajustados y texto con ajuste de línea.
#   - Las 5 columnas a llenar (etiqueta, confianza, es_relevante, nota,
#     frase_justificante) resaltadas en amarillo; el resto es solo lectura
#     visual (el script de fusión valida que `texto` no se haya editado).
#   - Validaciones desplegables: etiqueta ∈ {hawkish,dovish,neutral},
#     confianza ∈ {alta,media}, es_relevante ∈ {1,0}; frase_justificante
#     con aviso si supera 300 caracteres.
#   - Hoja "LEEME": guía mínima de llenado y devolución.
# Los saltos de línea internos de `texto` se aplanan a espacio (igual que
# preparar_gold_limpio.py); la verificación verbatim usa whitespace
# normalizado, así que el copiar-pegar desde la celda sigue siendo válido.
#
# DEVOLUCIÓN: el etiquetador devuelve el .xlsx llenado (a medias o completo)
# y se reintegra con:
#   python scripts/fusionar_gold_llenado.py data/muestras/gold_ciego_300.xlsx
# que también sigue aceptando el .csv limpio llenado.
#
# Ejecución:  python scripts/preparar_gold_xlsx.py
# =============================================================================

import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RUTA_MUESTRAS = Path(__file__).resolve().parent.parent / "data" / "muestras"
SRC = RUTA_MUESTRAS / "gold_ciego_300.csv"
DST = RUTA_MUESTRAS / "gold_ciego_300.xlsx"

COLS_LLENAR = ["etiqueta", "confianza", "es_relevante", "nota",
               "frase_justificante"]

ANCHOS = {
    "orden": 7, "intervencion_id": 24, "fecha_reunion": 13, "actor": 28,
    "cargo": 30, "texto": 100, "etiqueta": 12, "confianza": 12,
    "es_relevante": 13, "nota": 40, "frase_justificante": 60, "metodo": 13,
    "ronda": 12, "version_codebook": 13, "fecha": 12, "etiquetador": 12,
}

LEEME = [
    ("Gold ciego — cómo etiquetar (306 intervenciones)", None),
    ("", None),
    ("1. Trabaja SOLO en la hoja 'etiquetar'. No edites nada fuera de las "
     "5 columnas amarillas.", None),
    ("2. Por cada fila rellena:", None),
    ("   • etiqueta: hawkish / dovish / neutral (desplegable).", None),
    ("   • confianza: alta o media (desplegable).", None),
    ("   • es_relevante: 1 normal; 0 solo si es pura logística "
     "(apertura/cierre, ofrecer la palabra, agradecimientos, fijar fecha).",
     None),
    ("   • nota: breve; OBLIGATORIA si marcaste 0.", None),
    ("   • frase_justificante: copiar-pegar verbatim del texto (máx. 300 "
     "caracteres). Recomendada siempre; OBLIGATORIA si hawkish/dovish.", None),
    ("3. Si es_relevante=0, deja etiqueta=neutral y explica en nota.", None),
    ("4. Cada intervención se evalúa autónoma (sin usar recuerdos de ese "
     "actor) y SIN mirar las etiquetas de la IA: el ejercicio es ciego.", None),
    ("5. Ritmo sugerido: lotes de 30–40 por sesión. Puedes guardar a medias: "
     "el validador acepta filas pendientes.", None),
    ("6. Al terminar (o para revisión parcial), devuelve ESTE archivo .xlsx: "
     "el equipo lo reintegra al formato canónico con "
     "scripts/fusionar_gold_llenado.py.", None),
    ("", None),
    ("Criterio resumido (detalle en docs/codebook_v2.md):", None),
    ("• hawkish: favorece/justifica política MÁS restrictiva de lo que está "
     "sobre la mesa (subir TPM, apurar la normalización, mantener postura "
     "restrictiva cuando el menú apunta a relajar). El ancla es RELATIVA al "
     "menú del día.", None),
    ("• dovish: lo simétrico hacia lo expansivo (bajar, pausar un ciclo de "
     "alzas, mantener en el mínimo).", None),
    ("• neutral: describe sin tomar posición de dirección ('la inflación "
     "subió 0,6%' + nada más), presenta opciones SIN elegir, pregunta al "
     "staff, o reporta expectativas de terceros.", None),
    ("• Los votos explícitos mandan: 'voto por subir 25 pb' = hawkish.", None),
]


def aplanar(s: str) -> str:
    return s.replace("\r\n", "\n").replace("\r", "\n").replace("\n", " ")


def main() -> None:
    with open(SRC, newline="", encoding="utf-8") as f:
        filas = list(csv.reader(f))
    cabecera, datos = filas[0], filas[1:]
    assert all(len(r) == len(cabecera) for r in datos), "canónico mal formado"
    assert len(datos) == 306, f"se esperaban 306 filas, hay {len(datos)}"
    idx = {nombre: i for i, nombre in enumerate(cabecera)}

    wb = Workbook()
    ws = wb.active
    ws.title = "etiquetar"

    relleno_titulo = PatternFill("solid", fgColor="1F4E78")
    fuente_titulo = Font(bold=True, color="FFFFFF", size=11)
    relleno_editar = PatternFill("solid", fgColor="FFF2CC")
    fuente_cab_amarilla = Font(bold=True, color="7F6000", size=11)
    wrap = Alignment(wrap_text=True, vertical="top")
    centro = Alignment(horizontal="center", vertical="top")

    # --- cabecera ---
    for j, nombre in enumerate(cabecera, start=1):
        celda = ws.cell(row=1, column=j, value=nombre)
        if nombre in COLS_LLENAR:
            celda.fill = relleno_editar
            celda.font = fuente_cab_amarilla
        else:
            celda.fill = relleno_titulo
            celda.font = fuente_titulo
        celda.alignment = centro
        ws.column_dimensions[get_column_letter(j)].width = ANCHOS.get(
            nombre, 15)

    # --- datos ---
    for i, fila in enumerate(datos, start=2):
        for j, nombre in enumerate(cabecera, start=1):
            valor = aplanar(fila[idx[nombre]])
            if nombre == "orden":
                valor = int(valor)
            celda = ws.cell(row=i, column=j, value=valor)
            if nombre in ("texto", "nota", "frase_justificante"):
                celda.alignment = wrap
            elif nombre in ("orden", "etiqueta", "confianza", "es_relevante",
                            "fecha_reunion"):
                celda.alignment = centro
            else:
                celda.alignment = Alignment(vertical="top")
            if nombre in COLS_LLENAR:
                celda.fill = relleno_editar

    ultima = len(datos) + 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cabecera))}{ultima}"

    col = {n: get_column_letter(idx[n] + 1) for n in COLS_LLENAR}
    dv_etiqueta = DataValidation(type="list",
                                 formula1='"hawkish,dovish,neutral"',
                                 allow_blank=True,
                                 showErrorMessage=True,
                                 errorTitle="Etiqueta inválida",
                                 error="Usa: hawkish, dovish o neutral "
                                       "(minúsculas).")
    dv_etiqueta.add(f"{col['etiqueta']}2:{col['etiqueta']}{ultima}")
    dv_confianza = DataValidation(type="list", formula1='"alta,media"',
                                  allow_blank=True, showErrorMessage=True,
                                  errorTitle="Confianza inválida",
                                  error="Usa: alta o media.")
    dv_confianza.add(f"{col['confianza']}2:{col['confianza']}{ultima}")
    dv_relevante = DataValidation(type="list", formula1='"1,0"',
                                  allow_blank=True, showErrorMessage=True,
                                  errorTitle="Valor inválido",
                                  error="Usa 1 (normal) o 0 (pura logística).")
    dv_relevante.add(f"{col['es_relevante']}2:{col['es_relevante']}{ultima}")
    dv_frase = DataValidation(type="textLength", operator="lessThanOrEqual",
                              formula1="300", allow_blank=True,
                              showErrorMessage=True, errorStyle="warning",
                              errorTitle="Frase muy larga",
                              error="La frase supera 300 caracteres: "
                                    "recórtala a lo esencial.")
    dv_frase.add(
        f"{col['frase_justificante']}2:{col['frase_justificante']}{ultima}")
    for dv in (dv_etiqueta, dv_confianza, dv_relevante, dv_frase):
        ws.add_data_validation(dv)

    # --- hoja LEEME ---
    leeme = wb.create_sheet("LEEME")
    leeme.column_dimensions["A"].width = 130
    for i, (texto, _) in enumerate(LEEME, start=1):
        celda = leeme.cell(row=i, column=1, value=texto)
        celda.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            celda.font = Font(bold=True, size=13, color="1F4E78")

    wb.save(DST)

    # --- verificación: identidad con el canónico (whitespace normalizado) ---
    from openpyxl import load_workbook
    wb2 = load_workbook(DST, read_only=True, data_only=True)
    ws2 = wb2["etiquetar"]
    leidas = list(ws2.iter_rows(values_only=True))
    assert [str(c) for c in leidas[0]] == cabecera, "cabecera alterada"
    assert len(leidas) - 1 == 306, "cambió el número de filas"
    norm = lambda s: " ".join(str(s if s is not None else "").split())
    for f_o, f_x in zip(datos, leidas[1:]):
        assert str(f_x[idx["intervencion_id"]]) == f_o[idx["intervencion_id"]]
        for j, nombre in enumerate(cabecera):
            if nombre in COLS_LLENAR:
                continue
            assert norm(f_x[j]) == norm(f_o[idx[nombre]]), (
                f"contenido alterado: orden {f_o[0]}, columna {nombre}")

    print(f"OK: {DST.name}")
    print("  filas: 306 | columnas: 16 | hoja 'etiquetar' + hoja 'LEEME'")
    print("Abre ESTE archivo en Excel/LibreOffice para etiquetar.")


if __name__ == "__main__":
    sys.exit(main())
